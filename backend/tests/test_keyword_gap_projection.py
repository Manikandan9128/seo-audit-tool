from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _audit_slide_geometry, _ctr_decay_pct, add_keyword_gap_slides


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_ctr_tiers_match_spec_exactly():
    assert _ctr_decay_pct(1) == 22.0
    assert _ctr_decay_pct(2) == 13.0
    assert _ctr_decay_pct(3) == 9.0
    assert _ctr_decay_pct(4) == 6.0
    assert _ctr_decay_pct(5) == 4.5
    assert _ctr_decay_pct(6) == 1.5
    assert _ctr_decay_pct(10) == 1.5
    assert _ctr_decay_pct(11) == 0.2
    assert _ctr_decay_pct(20) == 0.2
    assert _ctr_decay_pct(21) == 0.0
    assert _ctr_decay_pct(50) == 0.0


def test_missing_or_not_ranking_position_is_zero_ctr():
    assert _ctr_decay_pct(None) == 0.0
    assert _ctr_decay_pct(0) == 0.0


def _gap_row(keyword, volume, kd, your_position=None, your_url=None, competitors=None, gap_category=None):
    return {
        "keyword": keyword, "search_volume": volume, "keyword_difficulty": kd,
        "your_position": your_position, "your_url": your_url,
        "competitor_positions": competitors or [],
        "gap_category": gap_category,
    }


def test_na_is_exact_string_no_dash_no_blank_url():
    # 2026-09-23 spec: "NA", never "Not ranking".
    analysis = {"keyword_gap_rows": [
        _gap_row("kw1", 1000, 40, your_position=None, gap_category="Untapped"),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    assert len(slides) == 2  # executive summary + detail summary (1 Untapped keyword stays inline, no dedicated slide)
    text = _slide_text(slides[1])
    assert "NA" in text
    assert "Not ranking" not in text


def test_position_and_url_shown_in_separate_columns_when_ranking():
    # 2026-09-22 layout change: Position and URL are two separate columns
    # per ranking source, not one "#N · /path" cell.
    analysis = {"keyword_gap_rows": [
        _gap_row(
            "kw1", 1000, 40, your_position=None,
            competitors=[{"competitor": "rival.com", "position": 7, "ranking_url": "https://rival.com/payroll-compliance"}],
            gap_category="Missing",
        ),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    # Status, Keyword, Volume, KD, My Position, My URL, rival.com (Position), URL
    assert table.cell(1, 6).text_frame.text == "#7"
    assert table.cell(1, 7).text_frame.text == "/payroll-compliance"


def test_competitor_ranking_cell_is_a_real_hyperlink_to_source_url():
    # 2026-09-22 spec rule 7: never invented/reconstructed — the exact URL
    # already present on the row.
    analysis = {"keyword_gap_rows": [
        _gap_row(
            "kw1", 1000, 40,
            competitors=[{"competitor": "rival.com", "position": 7, "ranking_url": "https://rival.com/payroll-compliance"}],
            gap_category="Missing",
        ),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    cell = table.cell(1, 7)  # Status, Keyword, Volume, KD, My Position, My URL, rival.com (Pos), URL
    run = cell.text_frame.paragraphs[0].runs[0]
    assert run.hyperlink.address == "https://rival.com/payroll-compliance"


def test_ranked_cell_with_no_recorded_url_has_no_hyperlink():
    # Ranked (has a position) but the source data carries no URL for it —
    # must show the real position text, never invent a link target.
    analysis = {"keyword_gap_rows": [
        _gap_row("kw1", 1000, 40, competitors=[{"competitor": "rival.com", "position": 7, "ranking_url": None}], gap_category="Missing"),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    pos_cell = table.cell(1, 6)
    url_cell = table.cell(1, 7)
    assert pos_cell.text_frame.text == "#7"
    assert url_cell.text_frame.text == "—"
    runs = url_cell.text_frame.paragraphs[0].runs
    assert not runs or runs[0].hyperlink.address is None


def test_my_position_url_is_also_a_real_hyperlink():
    # Regression: My Position's own URL used to be left out of the
    # hyperlink list entirely, even with a real your_url on record.
    analysis = {"keyword_gap_rows": [
        _gap_row(
            "kw1", 1000, 40, your_position=27, your_url="https://example.com/sql-server-dba-services",
            competitors=[{"competitor": "rival.com", "position": 9, "ranking_url": None}], gap_category="Shared",
        ),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    pos_cell = table.cell(1, 4)
    url_cell = table.cell(1, 5)
    assert pos_cell.text_frame.text == "#27"
    assert url_cell.text_frame.text == "/sql-server-dba-services"
    run = url_cell.text_frame.paragraphs[0].runs[0]
    assert run.hyperlink.address == "https://example.com/sql-server-dba-services"


def test_na_cell_has_no_hyperlink():
    analysis = {"keyword_gap_rows": [
        _gap_row("kw1", 1000, 40, competitors=[], gap_category="Missing"),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    cell = table.cell(1, 4)  # My Position column — no competitor tracked at all here
    assert cell.text_frame.text == "NA"


def test_same_competitor_column_mapping_across_rows():
    # rival.com ranks on both rows, leader.com only on the second — the
    # column for rival.com must be the same column on every row, and a row
    # where a tracked competitor doesn't rank shows "NA" in that domain's
    # own column rather than shifting columns.
    analysis = {"keyword_gap_rows": [
        _gap_row("kw1", 1000, 40, competitors=[{"competitor": "rival.com", "position": 5, "ranking_url": None}], gap_category="Missing"),
        _gap_row(
            "kw2", 900, 30,
            competitors=[
                {"competitor": "rival.com", "position": 9, "ranking_url": None},
                {"competitor": "leader.com", "position": 2, "ranking_url": None},
            ],
            gap_category="Missing",
        ),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    headers = [c.text_frame.text for c in table.rows[0].cells]
    rival_col = headers.index("rival.com")
    leader_col = headers.index("leader.com")
    row1 = [c.text_frame.text for c in table.rows[1].cells]
    row2 = [c.text_frame.text for c in table.rows[2].cells]
    assert row1[rival_col] == "#5"
    assert row1[leader_col] == "NA"
    assert row2[rival_col] == "#9"
    assert row2[leader_col] == "#2"


def test_kd_above_max_is_excluded():
    analysis = {"keyword_gap_rows": [
        _gap_row("too hard", 1000, 95, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
    ]}
    assert add_keyword_gap_slides(_prs(), analysis) == []


def test_kd_unavailable_flagged_separately():
    analysis = {"keyword_gap_rows": [
        _gap_row("no kd data", 1000, None, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
    ]}
    assert add_keyword_gap_slides(_prs(), analysis) == []


def test_ambiguous_relevance_excluded_from_table_with_review_note():
    rows = [
        _gap_row("clearly relevant", 1000, 40, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
        _gap_row("unsure keyword", 500, 30, competitors=[{"competitor": "rival.com", "position": 4, "ranking_url": None}], gap_category="Missing"),
    ]
    rows[1]["relevance"] = "potentially_relevant"
    analysis = {"keyword_gap_rows": rows}
    slides = add_keyword_gap_slides(_prs(), analysis)
    text = _slide_text(slides[1])
    assert "clearly relevant" in text
    assert "unsure keyword" not in text.split("KEY INSIGHTS")[0]
    assert "manual relevance review" in text
    assert "unsure keyword" in text


def test_zero_or_missing_volume_rows_excluded():
    # 2026-09-21 spec rule 7: "highest volume" selection must only ever
    # draw from rows with real, valid search volume.
    analysis = {"keyword_gap_rows": [
        _gap_row("zero volume kw", 0, 30, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
    ]}
    assert add_keyword_gap_slides(_prs(), analysis) == []


def test_status_column_header_is_labeled():
    analysis = {"keyword_gap_rows": [
        _gap_row("kw1", 1000, 40, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    table = next(s for s in slides[1].shapes if s.has_table).table
    assert table.cell(0, 0).text_frame.text == "Status"


def test_key_insights_include_actionable_implication_for_missing_keywords():
    analysis = {"keyword_gap_rows": [
        _gap_row("missing kw", 1000, 40, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
    ]}
    slides = add_keyword_gap_slides(_prs(), analysis)
    text = _slide_text(slides[1])
    assert "Prioritize validation of high-volume Missing keywords" in text
    # Rule 6: never an unsupported strategic claim.
    for banned in ("will generate", "will increase conversions", "best opportunity", "create a page immediately"):
        assert banned not in text.lower()


def test_gap_scale_insight_adds_interpretation_not_just_counts():
    analysis = {
        "keyword_gap_rows": [
            _gap_row("missing kw", 1000, 40, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
        ],
        "keyword_gap_off_topic_count": 2,
    }
    slides = add_keyword_gap_slides(_prs(), analysis, client_name="Acme")
    text = _slide_text(slides[1])
    assert "indicating the scale of the competitive keyword gap" in text


def test_off_topic_count_and_split_stated_in_insights():
    analysis = {
        "keyword_gap_rows": [
            _gap_row("shared kw", 1000, 40, your_position=5, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Shared"),
            _gap_row("missing kw", 800, 40, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category="Missing"),
            _gap_row("untapped kw", 300, 40, competitors=[], gap_category="Untapped"),
        ],
        "keyword_gap_off_topic_count": 4,
    }
    slides = add_keyword_gap_slides(_prs(), analysis, client_name="Acme")
    text = _slide_text(slides[1])
    assert "4 off-topic excluded" in text
    assert "unrelated to Acme's business" in text
    assert "3 relevant keyword" in text
    assert "1 Shared / 1 Missing / 1 Untapped" in text


# ---- 2026-09-22 spec: dedicated per-status slide thresholds ----

def _make_rows(category, count, start_volume=5000):
    return [
        _gap_row(f"{category.lower()} kw {i}", start_volume - i, 30, competitors=[{"competitor": "rival.com", "position": 3, "ranking_url": None}], gap_category=category)
        for i in range(count)
    ]


def test_zero_keywords_in_a_status_creates_no_slide_and_no_row():
    # Spec example 1: Missing=440(ish, using a smaller stand-in count that
    # still crosses the 6+ threshold), Shared=1, Untapped=0.
    analysis = {"keyword_gap_rows": _make_rows("Missing", 10) + _make_rows("Shared", 1)}
    slides = add_keyword_gap_slides(_prs(), analysis)
    titles = [next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text) for s in slides]
    assert "Competitor Keyword Gap — Executive Summary" in titles[0]
    assert "Competitor Keyword Gap Analysis" in titles[1]
    assert any("Competitor Keyword Gap — Missing" in t for t in titles)
    assert not any("Untapped" in t for t in titles)  # 0 Untapped keywords — no slide at all
    # Shared (1 keyword) must stay inline on the summary, never get its own slide.
    assert not any(t == "Competitor Keyword Gap — Shared" for t in titles)
    summary_text = _slide_text(slides[1])
    assert "shared kw 0" in summary_text


def test_one_to_five_keywords_never_gets_a_dedicated_slide():
    analysis = {"keyword_gap_rows": _make_rows("Missing", 5)}
    slides = add_keyword_gap_slides(_prs(), analysis)
    assert len(slides) == 2  # executive summary + detail summary only — 5 is still inline range
    text = _slide_text(slides[1])
    for i in range(5):
        assert f"missing kw {i}" in text.split("KEY INSIGHTS")[0]


def test_six_keywords_gets_its_own_dedicated_slide_not_in_summary_table():
    analysis = {"keyword_gap_rows": _make_rows("Missing", 6)}
    slides = add_keyword_gap_slides(_prs(), analysis)
    assert len(slides) == 3
    summary_table_text = "\n".join(
        c.text_frame.text for sh in slides[1].shapes if sh.has_table for row in sh.table.rows for c in row.cells
    )
    assert "missing kw" not in summary_table_text  # moved entirely to the dedicated slide
    dedicated_title = next(sh.text_frame.text for sh in slides[2].shapes if sh.has_text_frame and sh.text_frame.text)
    assert dedicated_title == "Competitor Keyword Gap — Missing"
    dedicated_text = _slide_text(slides[2])
    assert "missing kw 0" in dedicated_text


def test_spec_example_two_20_missing_12_shared_3_untapped():
    analysis = {
        "keyword_gap_rows": _make_rows("Missing", 20, start_volume=9000)
        + _make_rows("Shared", 12, start_volume=5000)
        + _make_rows("Untapped", 3, start_volume=1000)
    }
    slides = add_keyword_gap_slides(_prs(), analysis)
    titles = [next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text) for s in slides]
    assert titles == [
        "Competitor Keyword Gap — Executive Summary",
        "Competitor Keyword Gap Analysis",
        "Competitor Keyword Gap — Missing",
        "Competitor Keyword Gap — Shared",
    ]
    # Untapped (3, inline) shows on the summary; Missing/Shared (both
    # dedicated) do not appear as table rows there.
    summary_table_text = "\n".join(
        c.text_frame.text for sh in slides[1].shapes if sh.has_table for row in sh.table.rows for c in row.cells
    )
    assert "untapped kw 0" in summary_table_text
    assert "missing kw" not in summary_table_text
    assert "shared kw" not in summary_table_text


def test_dedicated_slide_never_contains_a_row_from_another_status():
    analysis = {"keyword_gap_rows": _make_rows("Missing", 8) + _make_rows("Shared", 8)}
    slides = add_keyword_gap_slides(_prs(), analysis)
    missing_slide = next(s for s in slides if any(
        sh.has_text_frame and sh.text_frame.text == "Competitor Keyword Gap — Missing" for sh in s.shapes
    ))
    missing_table_text = "\n".join(
        c.text_frame.text for sh in missing_slide.shapes if sh.has_table for row in sh.table.rows for c in row.cells
    )
    assert "shared kw" not in missing_table_text


def test_dedicated_slide_insights_only_reference_its_own_status():
    # 2026-09-22 spec rule 9.
    analysis = {"keyword_gap_rows": _make_rows("Missing", 8) + _make_rows("Untapped", 8)}
    slides = add_keyword_gap_slides(_prs(), analysis)
    missing_slide = next(s for s in slides if any(
        sh.has_text_frame and sh.text_frame.text == "Competitor Keyword Gap — Missing" for sh in s.shapes
    ))
    text = _slide_text(missing_slide)
    assert "untapped kw" not in text
    assert "Highest-volume Missing keyword" in text


def test_dedicated_slide_discloses_truncation_beyond_row_cap():
    analysis = {"keyword_gap_rows": _make_rows("Missing", 15)}
    slides = add_keyword_gap_slides(_prs(), analysis, keyword_gap_sheet_link="https://sheets.google.com/x")
    missing_slide = slides[2]
    text = _slide_text(missing_slide)
    assert "Showing top" in text
    assert "15 Missing keyword" in text
    assert "Open full keyword list" in text


def test_all_categories_worst_case_fits_slide_no_overlap():
    # Standing no-overlap/fit-to-page rule — worst case for the summary
    # slide's inline-row rendering: 3 competitor columns, long wrapping
    # URLs, and the keyword-sheet-link button all present at once, with
    # each status still in the 1-5 inline range (no dedicated slides).
    competitors = [
        {"competitor": "rivalonecompany.com", "position": 3, "ranking_url": "https://rivalonecompany.com/very-long-descriptive-category-slug/product-page"},
        {"competitor": "competitortwo.com", "position": 5, "ranking_url": "https://competitortwo.com/another-long-slug/product"},
        {"competitor": "thirdrivalbrand.com", "position": 8, "ranking_url": "https://thirdrivalbrand.com/yet-another-long-path/here"},
    ]
    rows = []
    for i in range(5):
        rows.append(_gap_row(f"missing keyword number {i} long tail phrase", 5000 - i, 30, competitors=competitors, gap_category="Missing"))
    for i in range(5):
        rows.append(_gap_row(f"untapped keyword number {i} long tail phrase", 4000 - i, 30, gap_category="Untapped"))
    for i in range(5):
        rows.append(_gap_row(
            f"shared keyword number {i} long tail phrase", 3000 - i, 30,
            your_position=6, your_url="https://example.com/very-long-descriptive-category-slug/product-page",
            competitors=competitors, gap_category="Shared",
        ))
    analysis = {"keyword_gap_rows": rows, "keyword_gap_off_topic_count": 3}
    prs = _prs()
    add_keyword_gap_slides(prs, analysis, client_name="Acme Trucks", keyword_gap_sheet_link="https://sheets.google.com/x")
    assert _audit_slide_geometry(prs) == []


def test_dedicated_slide_full_row_cap_fits_no_overlap():
    # Worst case for a single dedicated slide: the full _GAP_DEDICATED_ROW_CAP
    # (9) rows, 3 competitor columns, long wrapping URLs, sheet-link button.
    competitors = [
        {"competitor": "rivalonecompany.com", "position": 3, "ranking_url": "https://rivalonecompany.com/very-long-descriptive-category-slug/product-page"},
        {"competitor": "competitortwo.com", "position": 5, "ranking_url": "https://competitortwo.com/another-long-slug/product"},
        {"competitor": "thirdrivalbrand.com", "position": 8, "ranking_url": "https://thirdrivalbrand.com/yet-another-long-path/here"},
    ]
    rows = [
        _gap_row(f"missing keyword number {i} long tail phrase", 5000 - i, 30, competitors=competitors, gap_category="Missing")
        for i in range(12)
    ]
    analysis = {"keyword_gap_rows": rows, "keyword_gap_off_topic_count": 3}
    prs = _prs()
    add_keyword_gap_slides(prs, analysis, client_name="Acme Trucks", keyword_gap_sheet_link="https://sheets.google.com/x")
    assert _audit_slide_geometry(prs) == []

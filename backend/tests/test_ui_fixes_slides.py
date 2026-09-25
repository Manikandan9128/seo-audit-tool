"""UI-Level Fixes rebuild, Part B + Part C QA (2026-09-25 spec): the two
new slides (Top 5 Issues, Expected Outcomes) and the 0-issues edge case."""

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from app.reporting.pptx_builder import (
    DEFAULT_ACCENT, SLIDE_H, SLIDE_W, _accent, _audit_slide_geometry, _contrast_ratio, _mix_with_white,
    _readable_text_on, _theme, add_ui_fixes_no_issues_slide, add_ui_fixes_outcomes_slide, add_ui_fixes_top5_slide,
)


@pytest.fixture(autouse=True)
def _restore_theme_after_each_test():
    # _theme is thread-local but shared across every test in this whole
    # pytest process (single thread) — a test that sets a custom accent/
    # footer and never restores it would leak into whatever test runs
    # next, in this file or (worse) an unrelated one later in the suite.
    yield
    _theme["accent"] = DEFAULT_ACCENT
    _theme["footer"] = ""


def _prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _issue(n=1, **overrides):
    base = {
        "title": f"Issue {n} title", "evidence": f"Evidence {n} measured fact", "where": f"Location {n}",
        "fix": f"Fix {n} action", "outcome_tags": ["clarity"], "priority": "High", "device": "Both", "impact_score": 8,
    }
    base.update(overrides)
    return base


def _issues(n, **shared):
    return [_issue(i, where=f"Location {i}", **shared) for i in range(1, n + 1)]


def _reset_theme():
    _theme["accent"] = RGBColor(0xFF, 0x5C, 0x00)
    _theme["footer"] = "Geopits  ·  geopits.com"


# --- contrast guard helpers ---------------------------------------------------

def test_contrast_ratio_white_on_white_is_1():
    assert round(_contrast_ratio(RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xFF, 0xFF, 0xFF)), 2) == 1.0


def test_contrast_ratio_black_on_white_is_21():
    assert round(_contrast_ratio(RGBColor(0, 0, 0), RGBColor(0xFF, 0xFF, 0xFF)), 1) == 21.0


def test_readable_text_uses_white_on_a_dark_brand():
    assert _readable_text_on(RGBColor(0x1D, 0x4E, 0xD8)) == RGBColor(0xFF, 0xFF, 0xFF)


def test_readable_text_falls_back_to_dark_on_a_light_yellow_brand():
    assert _readable_text_on(RGBColor(0xF5, 0xE0, 0x42)) == RGBColor(0x14, 0x17, 0x1C)


def test_mix_with_white_at_zero_percent_is_white():
    assert _mix_with_white(RGBColor(0xFF, 0x5C, 0x00), 0.0) == RGBColor(0xFF, 0xFF, 0xFF)


def test_mix_with_white_at_full_percent_is_the_brand_color():
    brand = RGBColor(0xFF, 0x5C, 0x00)
    assert _mix_with_white(brand, 1.0) == brand


# --- Slide 1: Top 5 Issues -----------------------------------------------------

def test_top5_slide_returns_none_with_no_issues():
    _reset_theme()
    assert add_ui_fixes_top5_slide(_prs(), [], 0) is None


def test_top5_slide_title_says_top_5_when_5_or_more():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_top5_slide(prs, _issues(7), 7)
    text = _slide_text(slide)
    assert "UI-Level Fixes: Top 5 Issues" in text
    assert _audit_slide_geometry(prs) == []


def test_top5_slide_title_says_issues_found_when_fewer_than_5():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_top5_slide(prs, _issues(3), 3)
    text = _slide_text(slide)
    assert "UI-Level Fixes: Issues Found" in text
    assert "Top 5 Issues" not in text
    assert _audit_slide_geometry(prs) == []


def test_top5_slide_shows_only_the_top_5_even_with_more_validated():
    _reset_theme()
    prs = _prs()
    issues = _issues(20)
    slide = add_ui_fixes_top5_slide(prs, issues, 20)
    text = _slide_text(slide)
    assert "Location 1" in text and "Location 5" in text
    assert "Location 6" not in text
    assert _audit_slide_geometry(prs) == []


def test_top5_slide_shows_priority_and_device_chip_text():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_top5_slide(prs, [_issue(1, priority="High", device="Mobile")], 1)
    text = _slide_text(slide)
    assert "High" in text and "Mobile" in text
    assert _audit_slide_geometry(prs) == []


def test_top5_slide_includes_ga4_in_source_line_only_when_available():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_top5_slide(prs, _issues(2), 2, ga4_available=True)
    assert "+ GA4" in _slide_text(slide)

    prs2 = _prs()
    slide2 = add_ui_fixes_top5_slide(prs2, _issues(2), 2, ga4_available=False)
    assert "+ GA4" not in _slide_text(slide2)


def test_top5_slide_geometry_clean_at_1_through_5_issues():
    _reset_theme()
    for n in (1, 2, 3, 4, 5):
        prs = _prs()
        slide = add_ui_fixes_top5_slide(prs, _issues(n), n)
        issues = _audit_slide_geometry(prs)
        assert issues == [], f"{n} issues: {issues}"


def test_top5_slide_long_text_fields_do_not_overflow():
    _reset_theme()
    prs = _prs()
    long_issue = _issue(
        1, title="A very long issue title that describes a real usability problem in much more detail than usual",
        evidence="A long evidence sentence citing several measured facts from page_facts all at once, well past the normal length",
        where="A surprisingly long location description naming several nested sections of the page",
        fix="A long, detailed fix recommendation covering multiple concrete steps the team should take to resolve this",
    )
    slide = add_ui_fixes_top5_slide(prs, [long_issue] + _issues(4)[1:], 5)
    assert _audit_slide_geometry(prs) == []


# --- Slide 2: Expected Outcomes -------------------------------------------------

def test_outcomes_slide_returns_none_with_no_issues():
    _reset_theme()
    assert add_ui_fixes_outcomes_slide(_prs(), [], 0, {}, None, "Acme") is None


def test_outcomes_slide_shows_summary_counts():
    _reset_theme()
    prs = _prs()
    top5 = _issues(5, priority="High")
    slide = add_ui_fixes_outcomes_slide(prs, top5, 5, {"High": 5, "Medium": 0, "Low": 0}, None, "Acme")
    text = _slide_text(slide)
    assert "5" in text and "High" in text
    assert "All 5 shown on the previous slide" in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_shows_top5_of_n_and_button_when_more_than_5():
    _reset_theme()
    prs = _prs()
    top5 = _issues(5)
    slide = add_ui_fixes_outcomes_slide(prs, top5, 12, {"High": 12}, "https://sheet.example/x", "Acme")
    text = _slide_text(slide)
    assert "Top 5 of 12 shown on the previous slide" in text
    assert "Open full issue list (12)" in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_only_shows_cards_for_tags_present_on_top5():
    _reset_theme()
    prs = _prs()
    top5 = [
        _issue(1, where="A", outcome_tags=["clarity"]),
        _issue(2, where="B", outcome_tags=["cta"]),
    ]
    slide = add_ui_fixes_outcomes_slide(prs, top5, 2, {"High": 2}, None, "Acme")
    text = _slide_text(slide)
    assert "Clearer User Journey" in text
    assert "Higher CTA Engagement" in text
    assert "Lower Interaction Friction" not in text
    assert "Stronger Trust at Conversion Points" not in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_ignores_tags_on_issues_past_the_top5():
    # A tag that ONLY appears on issue #6+ must not get a card — the slide
    # only ever reflects what's shown on slide 1 (the top 5), never the
    # full validated list.
    _reset_theme()
    prs = _prs()
    top5 = _issues(5, outcome_tags=["clarity"])
    slide = add_ui_fixes_outcomes_slide(prs, top5, 20, {"High": 20}, "https://sheet.example/x", "Acme")
    text = _slide_text(slide)
    assert "Clearer User Journey" in text
    for absent in ("Higher CTA Engagement", "Lower Interaction Friction", "Stronger Trust at Conversion Points",
                   "Easier to Reach Us on Mobile", "Faster, More Focused Experience"):
        assert absent not in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_shows_issue_reference_numbers():
    _reset_theme()
    prs = _prs()
    top5 = [
        _issue(1, where="A", outcome_tags=["clarity"]),
        _issue(2, where="B", outcome_tags=["clarity"]),
        _issue(3, where="C", outcome_tags=["cta"]),
    ]
    slide = add_ui_fixes_outcomes_slide(prs, top5, 3, {"High": 3}, None, "Acme")
    text = _slide_text(slide)
    assert "Issues 1, 2" in text
    assert "Issue 3" in text


def test_outcomes_slide_measurement_section_adapts_to_ecommerce():
    _reset_theme()
    prs = _prs()
    top5 = _issues(3)
    slide = add_ui_fixes_outcomes_slide(
        prs, top5, 3, {"High": 3}, None, "Acme", business_type="ecommerce", primary_cta_text="Buy Now",
    )
    text = _slide_text(slide)
    assert "add-to-cart" in text
    assert "demo requests" not in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_flags_missing_heatmap_tool():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_outcomes_slide(prs, _issues(2), 2, {"High": 2}, None, "Acme", heatmap_tool_detected=False)
    text = _slide_text(slide)
    assert "Needs a tool such as Microsoft Clarity" in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_uses_ga4_baselines_when_available():
    _reset_theme()
    prs = _prs()
    ga4 = {"bounce_rate_pct": 52.3, "mobile_bounce_rate_pct": 61.0, "engagement_rate_pct": 47.5}
    slide = add_ui_fixes_outcomes_slide(prs, _issues(2), 2, {"High": 2}, None, "Acme", ga4_evidence=ga4)
    text = _slide_text(slide)
    assert "52%" in text and "61%" in text and "48%" in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_geometry_clean_with_all_6_outcome_tags():
    _reset_theme()
    prs = _prs()
    top5 = [_issue(i, where=f"L{i}", outcome_tags=[tag]) for i, tag in enumerate(
        ["clarity", "cta", "friction", "trust", "mobile_access", "speed_focus"], start=1
    )]
    slide = add_ui_fixes_outcomes_slide(prs, top5, 6, {"High": 6}, "https://sheet.example/x", "Acme")
    text = _slide_text(slide)
    for title in ("Clearer User Journey", "Higher CTA Engagement", "Lower Interaction Friction",
                  "Stronger Trust at Conversion Points", "Easier to Reach Us on Mobile", "Faster, More Focused Experience"):
        assert title in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_shows_measurement_section_even_with_all_6_outcome_tags():
    # Regression: all 6 outcome cards (2 rows x 3) used to push the
    # measurement section's 3rd content line past the footer — fixed by
    # tightening the outcome cards' own sizing rather than dropping the
    # measurement section (both are meant to always show together).
    _reset_theme()
    prs = _prs()
    top5 = [_issue(i, where=f"L{i}", outcome_tags=[tag]) for i, tag in enumerate(
        ["clarity", "cta", "friction", "trust", "mobile_access", "speed_focus"], start=1
    )]
    slide = add_ui_fixes_outcomes_slide(prs, top5, 6, {"High": 6}, "https://sheet.example/x", "Acme")
    text = _slide_text(slide)
    assert "HOW WE MEASURE THE IMPACT" in text
    assert "USER ENGAGEMENT" in text and "CONVERSION" in text and "QUALITATIVE FEEDBACK" in text
    assert _audit_slide_geometry(prs) == []


def test_outcomes_slide_light_brand_color_uses_dark_text_on_button():
    # Contrast guard end-to-end: a yellow brand must not put unreadable
    # white text on the full-list button.
    _theme["accent"] = RGBColor(0xF5, 0xE0, 0x42)
    _theme["footer"] = "Test Client  ·  test.com"
    prs = _prs()
    top5 = _issues(5)
    slide = add_ui_fixes_outcomes_slide(prs, top5, 12, {"High": 12}, "https://sheet.example/x", "Test Client")
    button = next(sh for sh in slide.shapes if sh.has_text_frame and "Open full issue list" in sh.text_frame.text)
    run_color = button.text_frame.paragraphs[0].runs[0].font.color.rgb
    assert run_color == RGBColor(0x14, 0x17, 0x1C)
    assert _audit_slide_geometry(prs) == []
    _reset_theme()


# --- Edge case: 0 issues ------------------------------------------------------

def test_no_issues_slide_lists_a_checklist():
    _reset_theme()
    prs = _prs()
    slide = add_ui_fixes_no_issues_slide(prs)
    text = _slide_text(slide)
    assert "No Major Issues Found" in text
    assert "Primary CTA visibility" in text
    assert _audit_slide_geometry(prs) == []

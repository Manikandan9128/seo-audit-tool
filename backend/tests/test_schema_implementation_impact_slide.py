from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _audit_slide_geometry, add_schema_implementation_impact_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def test_returns_none_when_no_impact():
    assert add_schema_implementation_impact_slide(_prs(), None) is None
    assert add_schema_implementation_impact_slide(_prs(), {"impact": []}) is None


def test_renders_with_two_points_clean_geometry():
    prs = _prs()
    impact = {"impact": [
        {"label": "Editorial Content", "text": "Clarifies the site's editorial content and its relationship to the topics covered."},
        {"label": "Product Pages", "text": "Makes key product attributes easier for search engines to interpret and associate with the offering."},
    ]}
    slide = add_schema_implementation_impact_slide(prs, impact)
    assert slide is not None
    assert _audit_slide_geometry(prs) == []


def test_renders_with_four_points_clean_geometry():
    prs = _prs()
    impact = {"impact": [
        {"label": "Editorial Content", "text": "Clarifies the site's editorial content and its relationship to the topics covered."},
        {"label": "Product Pages", "text": "Makes key product attributes easier for search engines to interpret and associate with the offering."},
        {"label": "Job Listings", "text": "Makes individual employment opportunities easier to interpret within applicable job-search experiences."},
        {"label": "Page Hierarchy", "text": "Establishes clearer relationships between pages and their position within the site's structure."},
    ]}
    slide = add_schema_implementation_impact_slide(prs, impact)
    assert slide is not None
    assert _audit_slide_geometry(prs) == []


def test_drops_a_point_with_a_directive_verb():
    impact = {"impact": [
        {"label": "Product Pages", "text": "Add Product structured data to the template."},
        {"label": "Page Hierarchy", "text": "Establishes clearer relationships between pages and their position within the site's structure."},
    ]}
    slide = add_schema_implementation_impact_slide(_prs(), impact)
    assert slide is not None
    texts = "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )
    assert "Add Product" not in texts
    assert "Page Hierarchy" in texts


def test_drops_a_point_with_a_number():
    impact = {"impact": [
        {"label": "Product Pages", "text": "Clarifies 12 product attributes for search engines."},
        {"label": "Page Hierarchy", "text": "Establishes clearer relationships between pages and their position within the site's structure."},
    ]}
    slide = add_schema_implementation_impact_slide(_prs(), impact)
    texts = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    assert "12 product" not in texts
    assert "Page Hierarchy" in texts


def test_returns_none_when_every_point_violates_rules():
    impact = {"impact": [{"label": "Product Pages", "text": "Add Product schema to 12 pages."}]}
    assert add_schema_implementation_impact_slide(_prs(), impact) is None

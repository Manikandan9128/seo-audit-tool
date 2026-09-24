"""Universal SEO Keyword Intelligence, Clustering & Targeting engine
(2026-09-23) — both scenarios: the AI/Semrush path (every keyword
clustered, cached AI calls) and the client's manual cluster sheet
(validated, never re-clustered)."""

import time
from unittest.mock import patch

from pptx import Presentation

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, _audit_slide_geometry, add_keyword_research_slide, add_strategic_keyword_clusters_slide,
)
from app.services.keyword_cluster_pipeline import _CAREER_ROUTE_CLUSTER_LABEL, build_final_keyword_clusters
from app.services.keyword_intelligence_service import (
    KeywordIntelligenceCache,
    annotate_keyword_rows,
    build_rule_groups,
    classify_with_cache,
    detect_intent,
    enrich_manual_clusters,
    gate_manual_rows,
    normalize_keyword,
    rule_group_name,
)
from app.services.keyword_relevance_service import (
    _rule_exclude, brand_token_variants, is_branded_or_near_brand, match_existing_page_for_cluster,
)
from app.services.semrush_analysis_service import analyze

_PHASE2_PATH = "app.services.keyword_cluster_pipeline.generate_phase2_candidate_clusters"
_PHASE3_PATH = "app.services.keyword_cluster_pipeline.generate_phase3_validated_clusters"
_THEMES_PATH = "app.services.keyword_cluster_pipeline.generate_business_themes"


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


# --- normalization / intent (§5, §8) --------------------------------------

def test_plural_and_word_order_share_one_core_key():
    assert normalize_keyword("tipper trucks")["core_key"] == normalize_keyword("truck tipper")["core_key"]
    assert normalize_keyword("certified payroll software")["core_key"] == normalize_keyword("certified payroll")["core_key"]


def test_geography_is_a_modifier_not_a_different_entity():
    assert normalize_keyword("bus price in india")["core_key"] == "bus"
    assert normalize_keyword("bus price in india")["geo"] == ["india"]


def test_intent_detection_uses_the_keywords_own_modifiers():
    assert detect_intent("how to hire construction workers")["intent"] == "Informational"
    assert detect_intent("truck price")["intent"] == "Transactional"
    assert detect_intent("union vs non union")["intent"] == "Comparison"
    assert detect_intent("truck dealer near me")["intent"] == "Local"
    assert detect_intent("bus mileage")["family"] == "Informational"
    # A bare entity term is only medium-confidence commercial, never certain.
    bare = detect_intent("tipper truck")
    assert bare["family"] == "Commercial" and bare["confidence"] < 70


# --- rule-based grouping (§17, §40, §41) ------------------------------------

def _grouped(keywords):
    rows = [{"keyword": k, "search_volume": 100} for k in keywords]
    annotate_keyword_rows(rows)
    return {rule_group_name(g): sorted(r["keyword"] for r in g["rows"]) for g in build_rule_groups(rows)}


def test_minor_modifier_variants_share_one_page():
    groups = _grouped(["certified payroll", "certified payrolls", "certified payroll software", "government certified payroll"])
    assert list(groups.values()) == [sorted(["certified payroll", "certified payrolls", "certified payroll software", "government certified payroll"])]


def test_informational_and_commercial_intent_get_separate_pages():
    groups = _grouped(["certified payroll", "what is certified payroll"])
    assert len(groups) == 2
    assert "Certified Payroll — Guides" in groups


def test_one_word_head_term_never_absorbs_different_products():
    groups = _grouped(["truck", "trucks", "tipper truck", "mining truck"])
    truck_group = next(v for v in groups.values() if "truck" in v)
    assert "tipper truck" not in truck_group and "mining truck" not in truck_group


# --- AI path: full coverage + cache (§28, no extra time) -------------------

def _many_rows(n):
    return [
        {"keyword": f"product{i} widget", "search_volume": 1000 - (i % 900), "intent": "Commercial", "page_category": "Landing Page"}
        for i in range(n)
    ]


def test_every_keyword_is_clustered_even_when_ai_fails():
    rows = _many_rows(1500)
    with patch(_THEMES_PATH, return_value={}), \
         patch(_PHASE2_PATH, return_value=({}, [])), \
         patch(_PHASE3_PATH, return_value=[]):
        build_final_keyword_clusters(rows, "Acme", None, None)
    assert all(r["cluster"] for r in rows), "no keyword may be left in an Ungrouped bucket"
    assert all(r["cluster_source"] == "rule" for r in rows)
    assert all("_core_tokens" not in r for r in rows), "internal keys must not leak into report data"


def test_ai_decision_on_a_representative_covers_its_whole_variant_group():
    rows = [
        {"keyword": "construction payroll", "search_volume": 900, "intent": "Commercial", "page_category": "Landing Page"},
        {"keyword": "construction payroll software", "search_volume": 500, "intent": "Commercial", "page_category": "Landing Page"},
        {"keyword": "construction payroll services", "search_volume": 300, "intent": "Commercial", "page_category": "Landing Page"},
    ]
    phase2 = ({"c1": {"main_entity": "payroll", "semantic_topic": "construction payroll", "justification": "", "keywords": ["construction payroll"]}}, [])
    phase3 = [{"cluster_name": "Construction Payroll Software", "primary_keyword": "construction payroll",
               "member_keywords": ["construction payroll"], "cluster_status": "Validated"}]
    with patch(_THEMES_PATH, return_value={}), patch(_PHASE2_PATH, return_value=phase2) as p2, \
         patch(_PHASE3_PATH, return_value=phase3):
        build_final_keyword_clusters(rows, "Acme", None, None)
    # One AI call, one representative in the prompt, three keywords placed.
    assert len(p2.call_args[0][0]) == 1
    assert p2.call_args[0][0][0]["group_size"] == 3
    assert {r["cluster"] for r in rows} == {"Construction Payroll Software"}
    assert all(r["cluster_source"] == "ai" for r in rows)
    assert sum(r["primary_or_secondary"] == "Primary" for r in rows) == 1


def test_regeneration_reuses_cached_ai_clusters_without_calling_ai():
    def run(cache):
        rows = [{"keyword": "construction payroll", "search_volume": 900, "intent": "Commercial", "page_category": "Landing Page"}]
        phase2 = ({"c1": {"keywords": ["construction payroll"]}}, [])
        phase3 = [{"cluster_name": "Construction Payroll", "primary_keyword": "construction payroll",
                   "member_keywords": ["construction payroll"], "cluster_status": "Validated"}]
        with patch(_THEMES_PATH, return_value={}), patch(_PHASE2_PATH, return_value=phase2) as p2, \
             patch(_PHASE3_PATH, return_value=phase3):
            build_final_keyword_clusters(rows, "Acme", "desc", None, cache=cache)
        return rows, p2.call_count

    cache = KeywordIntelligenceCache(None, "Acme", "desc")
    _rows, first_calls = run(cache)
    reloaded = KeywordIntelligenceCache(cache.data, "Acme", "desc")
    rows, second_calls = run(reloaded)
    assert first_calls == 1 and second_calls == 0
    assert rows[0]["cluster"] == "Construction Payroll"


def test_cache_is_dropped_when_business_context_changes():
    cache = KeywordIntelligenceCache(None, "Acme", "old description")
    cache.put_relevance({"widget": {"status": "Relevant", "reason": "x", "label": "highly_relevant"}})
    assert KeywordIntelligenceCache(cache.data, "Acme", "new description").get_relevance("widget") is None


def test_relevance_cache_only_sends_unjudged_keywords_to_the_ai():
    cache = KeywordIntelligenceCache(None, "Acme", None)
    cache.put_relevance({"known": {"status": "Relevant", "reason": "cached", "label": "highly_relevant"}})
    seen = []

    def fake_classify(name, domain, brands, keywords, desc):
        seen.append(list(keywords))
        return {k.lower(): {"status": "Relevant", "reason": "fresh", "label": "highly_relevant"} for k in keywords}

    result = classify_with_cache(fake_classify, cache, ["known", "new one"], "Acme", "acme.com", set(), None)
    assert seen == [["new one"]]
    assert result["known"]["reason"] == "cached" and result["new one"]["reason"] == "fresh"
    # Fail-open verdicts are never cached, so the next run retries them.
    cache.put_relevance({"flaky": {"status": "Unknown / Needs Review", "reason": "AI classification unavailable — needs manual review."}})
    assert cache.get_relevance("flaky") is None


def test_full_coverage_scales_without_adding_report_time():
    rows = [
        {"keyword": f"item{i} service{i % 50} pricing", "search_volume": 5000 - i, "intent": "Commercial"}
        for i in range(2000)
    ]
    pages = [{"page_url": f"https://acme.com/services/item{i}-service", "page_title": f"Item{i} Service"} for i in range(1400)]
    start = time.perf_counter()
    with patch(_THEMES_PATH, return_value={}), patch(_PHASE2_PATH, return_value=({}, [])), patch(_PHASE3_PATH, return_value=[]):
        build_final_keyword_clusters(rows, "Acme", None, pages)
    elapsed = time.perf_counter() - start
    assert all(r["cluster"] for r in rows)
    assert elapsed < 20, f"deterministic engine took {elapsed:.1f}s"


# --- confidence / action (§30, §53) ----------------------------------------

def test_every_cluster_gets_confidence_reason_and_spec_action():
    rows = [{"keyword": "construction payroll", "search_volume": 900, "intent": "Commercial", "page_category": "Landing Page"}]
    with patch(_THEMES_PATH, return_value={}), patch(_PHASE2_PATH, return_value=({}, [])), patch(_PHASE3_PATH, return_value=[]):
        build_final_keyword_clusters(rows, "Acme", None, None)
    r = rows[0]
    assert 0 <= r["cluster_confidence"] <= 85  # never High-certainty without SERP data
    assert r["cluster_confidence_level"] in ("High", "Medium", "Low")
    assert "no SERP data" in r["cluster_reason"]
    assert r["recommended_action"] in ("New URL Required", "Review — New URL Required")


# --- Target Keywords slide (AI path) ----------------------------------------

def test_routing_buckets_never_render_as_target_keywords():
    rows = [
        {"keyword": "construction payroll", "cluster": "Construction Payroll", "search_volume": 900},
        {"keyword": "construction payroll software", "cluster": "Construction Payroll", "search_volume": 400},
        {"keyword": "1099 contractor jobs", "cluster": _CAREER_ROUTE_CLUSTER_LABEL, "search_volume": 1000},
    ]
    slides = add_keyword_research_slide(_prs(), rows)
    text = "\n".join(_slide_text(s) for s in slides)
    assert "Target Keywords: Jobs / Careers" not in text
    assert "1099 contractor jobs" not in text
    assert "Not targeted: 1 job/career search(es)" in text


# --- Scenario A: manual cluster sheet (§21, §42, §45, §8) -------------------

def _gate(rows, relevance=None, competitor_domains=("ashokleyland.com",)):
    brands = set()
    for d in competitor_domains:
        brands |= brand_token_variants(d)
    return gate_manual_rows(
        rows, lambda kw: _rule_exclude(kw, set()), lambda kw: is_branded_or_near_brand(kw, brands), relevance,
    )


def test_manual_gate_removes_only_certain_cases_and_flags_ai_doubts():
    rows = [
        {"keyword": "school bus", "cluster": "Buses"},
        {"keyword": "indian bus xxx", "cluster": "Buses"},
        {"keyword": "ashok leyland bus price", "cluster": "Buses"},
        {"keyword": "peterbilt truck price in india", "cluster": "Truck Price"},
    ]
    relevance = {"peterbilt truck price in india": {"status": "Product/Service Mismatch", "reason": "Peterbilt is not sold here."}}
    kept, excluded, flagged = _gate(rows, relevance)
    # AI doubts never remove a client-sheet keyword (no missing data) — they
    # are flagged for client confirmation instead.
    assert [r["keyword"] for r in kept] == ["school bus", "peterbilt truck price in india"]
    reasons = {e["keyword"]: e["reason"] for e in excluded}
    assert "Adult" in reasons["indian bus xxx"]
    assert "brand" in reasons["ashok leyland bus price"]
    assert [f["keyword"] for f in flagged] == ["peterbilt truck price in india"]


def test_manual_gate_keeps_unjudged_keywords():
    kept, excluded, flagged = _gate([{"keyword": "tipper truck", "cluster": "Trucks"}], relevance={})
    assert len(kept) == 1 and not excluded and not flagged


def test_manual_cluster_enrichment_flags_intent_mismatch_and_split_without_rewriting_the_sheet():
    clusters = [{
        "cluster": "Buses",
        "keywords": [
            {"keyword": "school bus", "search_volume": 22200, "intent": "Commercial"},
            {"keyword": "bus mileage", "search_volume": 6600, "intent": "Commercial"},
            {"keyword": "truck sizes in india", "search_volume": 390, "intent": "Commercial"},
            {"keyword": "egg transport vehicle", "search_volume": 320, "intent": "Commercial"},
            {"keyword": "rv logistics tracking", "search_volume": 480, "intent": "Commercial"},
        ],
    }]
    pages = [{"page_url": "https://bharatbenz.com/buses/school-bus", "page_title": "School Bus"}]
    enrich_manual_clusters(clusters, pages, [{"keyword": "indian bus xxx", "cluster": "Buses", "reason": "Adult/explicit content"}],
                           match_existing_page_for_cluster)
    c = clusters[0]
    assert {m["keyword"] for m in c["intent_mismatches"]} == {"bus mileage", "truck sizes in india"}
    assert len(c["outliers"]) >= 2
    assert c["excluded"][0]["keyword"] == "indian bus xxx"
    assert c["confidence_level"] in ("High", "Medium", "Low") and c["recommended_action"]
    # Sheet values untouched.
    assert all(k["intent"] == "Commercial" for k in c["keywords"])
    assert all("_core_tokens" not in k for k in c["keywords"])

    slides = add_strategic_keyword_clusters_slide(_prs(), clusters)
    text = _slide_text(slides[0])
    assert "Intent corrected:" in text
    assert "Removed from this cluster:" in text
    assert "Confidence" in text
    assert "Every keyword here carries commercial" not in text


def test_validated_manual_slide_never_overlaps_worst_case():
    clusters = []
    for c in range(6):
        kws = [
            {"keyword": f"very long descriptive keyword phrase number {c} {k} mileage", "search_volume": 50000 - k,
             "keyword_difficulty": 40, "intent": "Commercial", "sub_category": f"Sub Topic {k % 3}"}
            for k in range(8)
        ]
        clusters.append({"cluster": f"Very Long Descriptive Cluster Topic Number {c}", "keywords": kws})
    excluded = [{"keyword": f"excluded keyword phrase {i}", "cluster": clusters[0]["cluster"], "reason": "Adult/explicit"} for i in range(5)]
    enrich_manual_clusters(clusters, [{"page_url": "https://x.com/a-very-long-path/that-goes-on", "page_title": "Phrase Number"}],
                           excluded, match_existing_page_for_cluster)
    prs = _prs()
    add_strategic_keyword_clusters_slide(prs, clusters)
    assert _audit_slide_geometry(prs) == []


# --- Competitor keyword gap brand filter (§21, §26, §46) --------------------

def test_keyword_gap_drops_competitor_brand_searches_but_keeps_comparisons():
    def row(kw, own, comp, vol):
        return {"keyword": kw, "search_volume": vol,
                "domain_positions": {"bharatbenz.com": own, "ashokleyland.com": comp, "mahindratruckandbus.com": 0}}
    records = [{
        "import_type": "keyword_gap", "is_own_site": True, "domain_label": None, "created_at": None,
        "parsed_data": {"rows": [
            row("truck price", 0, 5, 27100),
            row("ashok leyland price", 0, 2, 9900),
            row("mahindra", 0, 0, 368000),
            row("ashok leyland vs bharatbenz", 0, 3, 500),
        ]},
    }]
    result = analyze(records, own_domain="bharatbenz.com")
    kws = {r["keyword"] for r in result["keyword_gap_rows"]}
    assert "truck price" in kws and "ashok leyland vs bharatbenz" in kws
    assert "ashok leyland price" not in kws and "mahindra" not in kws
    assert result["keyword_gap_brand_excluded_count"] == 2


def test_rule_based_brand_verdicts_are_never_cached_across_callers():
    # The own-keyword filter passes the CLIENT's brand as brand_tokens, so a
    # rule marks "acme truck" as a brand search there — that verdict must not
    # leak into the manual-sheet check, which judges own-brand keywords by AI.
    cache = KeywordIntelligenceCache(None, "Acme", None)

    def fake_classify(name, domain, brands, keywords, desc):
        return {k.lower(): {"status": "Competitor Brand Search", "reason": "rule", "label": "exclude"} for k in keywords}

    classify_with_cache(fake_classify, cache, ["acme truck"], "Acme", "acme.com", {"acme"}, None)
    assert cache.get_relevance("acme truck") is None


def test_strict_competitor_brand_match_never_fires_on_similar_words():
    from app.services.keyword_relevance_service import is_competitor_brand_query
    brands = brand_token_variants("tatamotors.com") | brand_token_variants("ashokleyland.com") | brand_token_variants("vecv.in")
    assert is_competitor_brand_query("tata motors truck price", brands)
    assert is_competitor_brand_query("ashok leyland dost price", brands)
    assert not is_competitor_brand_query("taxi price", brands)
    assert not is_competitor_brand_query("data truck tracking", brands)
    assert not is_competitor_brand_query("deck truck", brands)


def test_page_matching_ignores_site_wide_boilerplate_words():
    # Every title carries the brand + "india" + "trucks" — those words must
    # not decide which page "covers" a cluster.
    pages = [{"page_url": f"https://bb.com/page-{i}", "page_title": f"BharatBenz Trucks India | Topic {i}"} for i in range(20)]
    pages.append({"page_url": "https://bb.com/important-information-for-customers", "page_title": "BharatBenz Trucks India Bus Information"})
    pages.append({"page_url": "https://bb.com/buses/school-bus", "page_title": "School Bus"})
    match = match_existing_page_for_cluster(["bus price in india", "school bus"], pages, "Landing Page")
    assert match["url"].endswith("/buses/school-bus")
    assert match_existing_page_for_cluster(["trucks in india"], pages, "Landing Page") is None


def test_manual_slide_shows_relevance_flags():
    clusters = [{"cluster": "Truck Price", "keywords": [
        {"keyword": "petrol truck price", "search_volume": 140, "intent": "Commercial"},
        {"keyword": "brabus price in india", "search_volume": 2400, "intent": "Commercial"},
        {"keyword": "chassis price", "search_volume": 480, "intent": "Commercial"},
    ]}]
    enrich_manual_clusters(clusters, None, [], match_existing_page_for_cluster,
                           flagged=[{"keyword": "brabus price in india", "cluster": "Truck Price", "reason": "Competitor Brand Search"}])
    text = _slide_text(add_strategic_keyword_clusters_slide(_prs(), clusters)[0])
    assert "brabus price in india" in text  # still shown — never silently dropped
    assert "Relevance check" in text and "confirm with the client" in text
    assert "brabus price in india †" not in text  # keyword cell stays exactly as the sheet has it
    assert 'Highest demand: "brabus' not in text  # demand lines never name a flagged keyword


# --- 2026-09-24 spec-gap fixes (§3, §8, §19-23, §30, §33, §55) --------------

def test_error_and_non_200_pages_are_never_keyword_targets():
    from app.services.keyword_relevance_service import is_live_target_page
    assert not is_live_target_page({"page_url": "https://bb.com/404.php", "http_status_code": "200"})
    assert not is_live_target_page({"page_url": "https://bb.com/trucks", "http_status_code": "301"})
    assert not is_live_target_page({"page_url": "https://bb.com/x", "page_title": "Page Not Found"})
    assert is_live_target_page({"page_url": "https://bb.com/error-codes-guide", "page_title": "Truck Error Codes"})
    pages = [
        {"page_url": "https://bb.com/404.php", "page_title": "Trucks Specifications Not Found"},
        {"page_url": "https://bb.com/about-us", "page_title": "About"},
    ]
    assert match_existing_page_for_cluster(["truck specifications"], pages, "Landing Page") is None


def test_manual_target_prefers_page_already_ranking():
    clusters = [{"cluster": "Parts", "keywords": [
        {"keyword": "truck parts", "search_volume": 880, "intent": "Commercial"},
        {"keyword": "truck spare parts", "search_volume": 880, "intent": "Commercial"},
    ]}]
    pages = [
        {"page_url": "https://bb.com/truck-parts-news", "page_title": "Truck Parts Truck Spare Parts"},
        {"page_url": "https://bb.com/genuine-parts", "page_title": "Genuine Parts"},
    ]
    enrich_manual_clusters(
        clusters, pages, [], match_existing_page_for_cluster,
        keyword_ranking={"truck parts": (6, "https://bb.com/genuine-parts")},
    )
    c = clusters[0]
    assert c["target_url"] == "https://bb.com/genuine-parts"
    assert c["match_strength"] == "strong" and c["ranking_evidence"]["position"] == 6
    text = _slide_text(add_strategic_keyword_clusters_slide(_prs(), clusters)[0])
    assert "already ranks #6" in text


def test_ranking_page_on_dead_url_is_ignored():
    from app.services.keyword_intelligence_service import ranking_page_target
    rows = [{"keyword": "a", "current_position": 3, "current_url": "https://www.bb.com/404.php/", "search_volume": 10}]
    assert ranking_page_target(rows, dead_urls={"https://bb.com/404.php"}) is None
    assert ranking_page_target([{"keyword": "a", "current_position": 35, "current_url": "https://bb.com/p"}]) is None


def test_manual_table_picks_verified_keywords_before_flagged_high_volume():
    from app.services.strategic_keyword_selection_service import select_strategic_clusters
    models = ["torres", "prima", "signa", "ultra", "boss", "dost", "partner", "bison"]
    rows = [{"keyword": f"{m} lorry price", "cluster": "Price", "search_volume": 100 + i} for i, m in enumerate(models)]
    rows += [
        {"keyword": "brabus price in india", "cluster": "Price", "search_volume": 2400, "relevance_status": "Competitor Brand Search"},
        {"keyword": "peterbilt truck price", "cluster": "Price", "search_volume": 390, "relevance_status": "Product/Service Mismatch"},
    ]
    chosen = [k["keyword"] for k in select_strategic_clusters(rows)[0]["keywords"]]
    assert "brabus price in india" not in chosen and len(chosen) == 8
    # With too few verified keywords, flagged ones fill the table — last.
    few = rows[:3] + rows[8:]
    chosen = [k["keyword"] for k in select_strategic_clusters(few)[0]["keywords"]]
    assert chosen[-2:] == ["brabus price in india", "peterbilt truck price"]


def test_manual_table_shows_corrected_intent_and_primary_keyword():
    clusters = [{"cluster": "Buses", "keywords": [
        {"keyword": "school bus", "search_volume": 22200, "intent": "Commercial"},
        {"keyword": "bus mileage", "search_volume": 6600, "intent": "Commercial"},
        {"keyword": "school bus price", "search_volume": 900, "intent": "Commercial"},
    ]}]
    enrich_manual_clusters(clusters, None, [], match_existing_page_for_cluster)
    c = clusters[0]
    assert c["primary_keyword"] == "school bus" and c["user_need"]
    by_kw = {k["keyword"]: k for k in c["keywords"]}
    assert by_kw["bus mileage"]["display_intent"] == "Informational"
    assert by_kw["bus mileage"]["intent"] == "Commercial"  # sheet value untouched
    slide = add_strategic_keyword_clusters_slide(_prs(), clusters)[0]
    table = next(sh.table for sh in slide.shapes if sh.has_table)
    cells = {row.cells[1].text: row.cells[4].text for row in table.rows}
    assert cells["bus mileage"] == "Commercial"  # table shows the sheet's intent; correction is an insight
    assert "Intent corrected:" in _slide_text(slide)
    assert 'Primary keyword: "school bus"' in _slide_text(slide)


def test_flagged_keywords_lower_cluster_confidence():
    from app.services.keyword_intelligence_service import score_cluster
    clean = [{"keyword": f"k{i}", "intent_family": "Commercial", "relevance_status": "Relevant"} for i in range(8)]
    doubtful = [dict(r) for r in clean]
    for r in doubtful[:3]:
        r["relevance_status"] = "Product/Service Mismatch"
    assert score_cluster(doubtful, "manual")["score"] < score_cluster(clean, "manual")["score"] - 10


def test_other_website_searches_are_called_out_separately():
    clusters = [{"cluster": "Trucks", "keywords": [
        {"keyword": "truck", "search_volume": 110000, "intent": "Commercial"},
        {"keyword": "truckspoint", "search_volume": 2400, "intent": "Commercial"},
    ]}]
    enrich_manual_clusters(clusters, None, [], match_existing_page_for_cluster, flagged=[
        {"keyword": "truckspoint", "cluster": "Trucks", "reason": "Other Website Search: a listing portal"},
    ])
    text = _slide_text(add_strategic_keyword_clusters_slide(_prs(), clusters)[0])
    assert '"truckspoint" are searches for another website' in text


def test_site_entity_summary_reads_own_sections_from_crawl():
    from app.services.keyword_relevance_service import site_entity_summary
    pages = [{"page_url": f"https://bb.com/select-my-truck/{n}-{i}", "http_status_code": "200"}
             for i, n in enumerate(["tipper", "tanker", "cement-mixer", "container"])]
    pages.append({"page_url": "https://bb.com/404.php"})
    summary = site_entity_summary(pages)
    assert "select my truck" in summary and "cement mixer" in summary and "404" not in summary
    assert site_entity_summary([{"page_url": "https://bb.com/about"}]) is None


def test_pipeline_existing_page_uses_ranking_url():
    from app.services.keyword_cluster_pipeline import _apply_existing_page_matching
    rows = [
        {"keyword": "tipper truck", "cluster": "Tippers", "current_position": 8,
         "current_url": "https://bb.com/select-my-truck/tipper-5", "search_volume": 14800},
        {"keyword": "tipper truck price", "cluster": "Tippers", "search_volume": 900},
    ]
    pages = [{"page_url": "https://bb.com/tipper-truck-news", "page_title": "Tipper Truck Price News"}]
    _apply_existing_page_matching(rows, pages)
    assert rows[0]["existing_page_url"] == "https://bb.com/select-my-truck/tipper-5"
    assert rows[1]["existing_page_match_strength"] == "strong"

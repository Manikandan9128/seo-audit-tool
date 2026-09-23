from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_conversion_seo_next_steps_slide, build_conversion_next_steps


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


def test_states_tracking_data_required_when_no_evidence_at_all():
    # Universal SEO Audit Engine spec (2026-09-20) section 39: no manual UX
    # pass and no backlink/funnel data must state "Tracking data required"
    # explicitly, not silently omit the slide.
    slide = add_conversion_seo_next_steps_slide(_prs(), None, 0)
    assert slide is not None
    text = _slide_text(slide)
    assert "Tracking data required" in text


def test_backlink_count_alone_is_not_conversion_evidence():
    # Conversion SEO spec 2026-09-23: backlink volume says nothing about the
    # conversion path, so it no longer produces a bullet.
    text = _slide_text(add_conversion_seo_next_steps_slide(_prs(), None, 500))
    assert "Tracking data required" in text and "500" not in text


def test_no_key_events_never_claims_poor_conversion():
    ev = {"events": [{"name": "page_view", "count": 9000, "key_events": 0, "kind": None}],
          "organic_landing_pages": [{"path": "/blog/payroll-guide", "sessions": 800, "key_events": 0, "session_key_event_rate": 0}],
          "key_events_configured": False}
    intro, items = build_conversion_next_steps(ev)
    joined = (intro + " ".join(items)).lower()
    assert "poor conversion" not in joined and "%" not in joined
    assert any("/blog/payroll-guide" in i and "product or service page" in i for i in items)


def test_low_rate_page_flagged_against_session_weighted_ga4_rate():
    ev = {"events": [{"name": "generate_lead", "count": 60, "key_events": 60, "kind": "conversion"}],
          "organic_landing_pages": [
              {"path": "/products/payroll", "sessions": 1000, "key_events": 50, "session_key_event_rate": 0.05},
              {"path": "/pricing", "sessions": 400, "key_events": 2, "session_key_event_rate": 0.005},
          ],
          "key_events_configured": True}
    intro, items = build_conversion_next_steps(ev)
    assert "generate_lead" in intro
    assert any(i.startswith("/pricing — 400 organic sessions with a 0.5% session key-event rate") for i in items)
    assert any("/products/payroll drives the most organic key events" in i for i in items)


def test_form_dropoff_only_when_start_and_submit_both_measured():
    ev = {"events": [{"name": "form_start", "count": 300, "key_events": 0, "kind": "path_signal"}], "key_events_configured": False}
    _intro, items = build_conversion_next_steps(ev)
    assert not any("Forms are started" in i for i in items)
    ev["events"].append({"name": "form_submit", "count": 40, "key_events": 0, "kind": "conversion"})
    intro, items = build_conversion_next_steps(ev)
    assert any("started 300 times but completed 40 times" in i for i in items)
    assert "not marked as key events" in intro


def test_classify_ga4_event():
    from app.services.ga4_service import classify_ga4_event
    assert classify_ga4_event("anything", 5) == "conversion"
    assert classify_ga4_event("form_submit", 0) == "conversion"
    assert classify_ga4_event("form_start", 0) == "path_signal"
    assert classify_ga4_event("scroll", 0) == "engagement"
    assert classify_ga4_event("page_view", 0) is None

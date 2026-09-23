from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_content_seo_next_steps_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
    return "\n".join(parts)


def test_uses_pipeline_existing_page_action_when_present():
    # 2026-09-20 spec section 47: the renderer must not re-derive its own
    # update-vs-create decision — it must use existing_page_action, already
    # computed upstream by keyword_cluster_pipeline.
    rows = [
        {"keyword": "construction payroll", "search_volume": 900, "intent": "Transactional", "cluster": "Construction Payroll",
         "existing_page_url": "https://example.com/payroll", "existing_page_action": "Optimize Existing Page"},
    ]
    slide = add_content_seo_next_steps_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "optimize the existing page (https://example.com/payroll)" in text


def test_falls_back_to_heuristic_when_pipeline_action_absent():
    rows = [
        {"keyword": "construction payroll", "search_volume": 900, "intent": "Transactional", "cluster": "Construction Payroll"},
    ]
    slide = add_content_seo_next_steps_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "create a new page — no existing page covers this topic closely enough" in text


def test_cannibalization_note_shown_once_per_url_not_per_cluster():
    rows = [
        {"keyword": "construction payroll", "search_volume": 900, "intent": "Transactional", "cluster": "Construction Payroll",
         "existing_page_url": "https://example.com/payroll", "existing_page_action": "Consolidate — Primary URL for this topic (https://example.com/payroll)",
         "cannibalization_status": "Potential cannibalization: https://example.com/payroll is also the best existing-page match for 1 other cluster(s) (Certified Payroll) — consolidate onto one page, differentiate the content, or confirm this is the right existing URL for this cluster."},
        {"keyword": "certified payroll", "search_volume": 200, "intent": "Transactional", "cluster": "Certified Payroll",
         "existing_page_url": "https://example.com/payroll", "existing_page_action": 'Differentiate or Redirect / Merge into "Construction Payroll"',
         "cannibalization_status": "Potential cannibalization: https://example.com/payroll is also the best existing-page match for 1 other cluster(s) (Construction Payroll) — consolidate onto one page, differentiate the content, or confirm this is the right existing URL for this cluster."},
    ]
    slide = add_content_seo_next_steps_slide(_prs(), rows)
    text = _slide_text(slide)
    assert text.count("Potential cannibalization") == 1


# --- Content SEO spec 2026-09-23 -------------------------------------------

def test_routing_buckets_and_competitor_rows_never_become_recommendations():
    rows = [
        {"keyword": "payroll software", "search_volume": 900, "intent": "Commercial", "cluster": "Payroll Software",
         "existing_page_action": "Create New Page"},
        {"keyword": "acme vs gusto", "search_volume": 800, "cluster": "Competitor / Comparison Opportunities",
         "competitor_status": "Competitor Comparison Opportunity"},
        {"keyword": "payroll jobs", "search_volume": 700, "cluster": "Jobs / Careers"},
        {"keyword": "odd query", "search_volume": 600, "cluster": "Needs Review — Relevance Unconfirmed"},
        {"keyword": "gusto login", "search_volume": 500, "cluster": "Gusto", "relevance_status": "Competitor Brand Search"},
    ]
    text = _slide_text(add_content_seo_next_steps_slide(_prs(), rows))
    assert "Payroll Software" in text
    for banned in ("Competitor / Comparison", "Jobs / Careers", "Needs Review", "gusto", "Gusto"):
        assert banned not in text


def test_cluster_without_evidence_gets_no_bullet_and_no_generic_filler():
    rows = [{"keyword": "payroll widget", "intent": "Commercial", "cluster": "Payroll Widget"}]
    assert add_content_seo_next_steps_slide(_prs(), rows) is None


def test_evidence_includes_ranking_and_gsc_when_supplied():
    rows = [{"keyword": "payroll software", "search_volume": 900, "position": 14, "gsc_impressions": 1200, "gsc_clicks": 30,
             "intent": "Commercial", "cluster": "Payroll Software", "existing_page_action": "Create New Page"}]
    text = _slide_text(add_content_seo_next_steps_slide(_prs(), rows))
    assert "ranking #14" in text and "1,200 Search Console impressions" in text


def test_existing_page_action_without_url_is_treated_as_new_page():
    rows = [{"keyword": "payroll software", "search_volume": 900, "intent": "Commercial", "cluster": "Payroll Software",
             "existing_page_action": "Optimize Existing Page", "existing_page_url": None}]
    text = _slide_text(add_content_seo_next_steps_slide(_prs(), rows))
    assert "None" not in text and "create a new page" in text


def test_duplicate_slashes_in_existing_url_are_collapsed():
    rows = [{"keyword": "payroll software", "search_volume": 900, "intent": "Commercial", "cluster": "Payroll Software",
             "existing_page_action": "Optimize Existing Page", "existing_page_url": "https://example.com///payroll"}]
    text = _slide_text(add_content_seo_next_steps_slide(_prs(), rows))
    assert "https://example.com/payroll" in text and "///" not in text

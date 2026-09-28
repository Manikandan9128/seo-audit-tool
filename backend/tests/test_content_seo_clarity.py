"""Content SEO Next Steps clarity (2026-09-28, from a Geopits deck review):
no conflicting keyword totals, no redirect advice for a page already in the
top 3, no pagination "cannibalization", topic totals matching Target
Keywords, and plain client-facing wording."""

from pptx import Presentation

from app.reporting import pptx_builder as b
from app.services import keyword_site_model
from app.services.keyword_strategy_depth import cannibalization_similarity


def _prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = b.SLIDE_W, b.SLIDE_H
    return prs


def _text(slide) -> str:
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def _row(keyword, volume, cluster, **extra):
    return {"keyword": keyword, "search_volume": volume, "intent": "Informational", "cluster": cluster, **extra}


def test_format_line_names_top_searches_not_a_competing_total():
    rows = [_row("how to fix slow mysql queries", 8100, "Tuning"), _row("database news", 50, "Tuning"),
            _row("mysql index tips", 900, "Tuning")]
    text = _text(b.add_content_seo_next_steps_slide(_prs(), rows))
    assert '"how to fix slow mysql queries" (8,100/month) and "mysql index tips" (900/month)' in text
    assert "relevant keyword(s) call for" not in text and "database news" not in text


def test_top_three_topic_is_never_told_to_redirect_or_merge():
    rows = [_row("remote dba services", 260, "Remote support", position=1,
                 existing_page_action='Differentiate or Redirect / Merge into "SQL server support"')]
    text = _text(b.add_content_seo_next_steps_slide(_prs(), rows))
    assert "protect the current ranking" in text and "Best current ranking: #1" in text and "Keep this page" in text
    assert "Redirect / Merge" not in text and "merge it into" not in text


def test_lower_ranked_overlapping_topic_gets_a_plain_choice():
    rows = [_row("remote dba services", 260, "Remote support", position=40,
                 existing_page_action='Differentiate or Redirect / Merge into "SQL server support"')]
    text = _text(b.add_content_seo_next_steps_slide(_prs(), rows))
    assert "separate it from a related topic" in text and "or merge it into that topic's page" in text


def test_priority_reads_as_plain_words():
    rows = [_row("database management services", 590, "Managed services", roadmap_priority="High")]
    text = _text(b.add_content_seo_next_steps_slide(_prs(), rows))
    assert "High priority: Managed services — " in text and "[High priority]" not in text


def test_topic_totals_match_target_keywords_when_rendered_in_one_deck():
    b._reset_deck_insights(active=True)
    try:
        b._register_cluster_totals("Managed services", 8, 1320)
        rows = [_row("database management services", 590, "Managed services"),
                _row("dba managed services", 870, "Managed services")]
        text = _text(b.add_content_seo_next_steps_slide(_prs(), rows))
    finally:
        b._reset_deck_insights(active=False)
    assert "8 keywords, 1,320 searches a month" in text


def test_paginated_listing_pages_are_not_duplicate_titles():
    def page(url):
        return {"url": url, "page_type": "blog", "primary_entity": "Geopits blog database insights sql server",
                "traffic_clicks": 0, "crawl_depth": 1}
    pages = [page(f"https://x.com/blog?8d31a6fe_page={n}") for n in range(1, 5)] + [page("https://x.com/blog/page/2")]
    assert keyword_site_model._duplicate_titles(pages) == []


def test_duplicate_title_bullet_is_readable_and_capped():
    site_model = {"duplicate_titles": [
        {"preferred_url": "https://x.com/a", "other_url": f"https://x.com/a-{n}", "similarity": 0.95, "page_type": "other"}
        for n in range(5)
    ]}
    strategy = {"cannibalization": cannibalization_similarity([], site_model)}
    text = _text(b.add_content_seo_next_steps_slide(_prs(), [_row("kw", 100, "Topic")], strategy))
    assert "6 pages have near-identical titles" in text and "(+2 more)" in text
    assert "(near-identical page titles)" not in text and "other other" not in text

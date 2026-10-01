from app.reporting.pptx_builder import _content_seo_eligible_rows
from app.services.technical_seo_service import _commerce_signals_present, aggregate_schema_validation


def _page(url, html_commerce, types=("Organization", "WebSite")):
    meta = {"schema_types_found": list(types), "schema_field_issues": []}
    if html_commerce is not None:
        meta["commerce_signals"] = html_commerce
    return {"url": url, "meta": meta}


def _bucket_pages(result, page_type):
    return next((b["pages"] for b in result["by_page_type"] if b["page_type"] == page_type), 0)


def test_commerce_signal_detection():
    assert _commerce_signals_present('<button>Add to Cart</button>')
    assert _commerce_signals_present('<span class="price">$5</span>')
    assert not _commerce_signals_present('<a href="/demo">Book a demo</a><div class="pricing-hero">')


def test_saas_product_url_is_not_a_product_page():
    pages = [_page("https://x.com/product/payroll", False), _page("https://x.com/product/hr", False)]
    result = aggregate_schema_validation(pages)
    assert _bucket_pages(result, "Product") == 0
    assert _bucket_pages(result, "Other Pages") == 2


def test_shop_product_url_with_price_stays_product():
    result = aggregate_schema_validation([_page("https://x.com/product/shoe", True)])
    assert _bucket_pages(result, "Product") == 1


def test_old_crawl_job_without_commerce_field_keeps_url_behaviour():
    result = aggregate_schema_validation([_page("https://x.com/product/payroll", None)])
    assert _bucket_pages(result, "Product") == 1


def test_competitor_only_adjacent_keyword_is_dropped_but_ranking_one_kept():
    rows = [
        {"keyword": "employee benefits", "search_volume": 201000, "cluster": "HR", "source": "Keyword Gap",
         "relevance_status": "Adjacent / Potential"},
        {"keyword": "minimum wage memphis", "search_volume": 70, "cluster": "HR", "source": "Keyword Gap, Organic Positions",
         "relevance_status": "Adjacent / Potential", "position": 25},
        {"keyword": "construction payroll software", "search_volume": 1300, "cluster": "Payroll", "source": "Keyword Gap",
         "relevance_status": "Core Relevant"},
    ]
    kept = [r["keyword"] for r in _content_seo_eligible_rows(rows)]
    assert kept == ["minimum wage memphis", "construction payroll software"]


def test_guard_never_empties_the_slide():
    rows = [{"keyword": "a b", "search_volume": 10, "cluster": "X", "source": "Keyword Gap", "relevance_status": "Adjacent / Potential"}]
    assert len(_content_seo_eligible_rows(rows)) == 1


def test_client_safe_insight_strips_internal_scaffolding():
    from app.services.structured_data_insights_service import client_safe_insight
    raw = ("ISSUE: FAQPage status unconfirmed. EVIDENCE: flagged with ELIGIBILITY_CHECK_STALE, "
           "so not verified. ACTION: Fix: re-validate.")
    out = client_safe_insight(raw)
    for banned in ("ISSUE:", "EVIDENCE:", "ACTION:", "ELIGIBILITY_CHECK_STALE"):
        assert banned not in out
    assert "re-validate" in out


def test_programmatic_slide_has_no_internal_wording():
    from pptx import Presentation
    from app.reporting.pptx_builder import add_programmatic_seo_slide
    rows = [{"keyword": k, "search_volume": v, "cluster": "Federal Minimum Wage"} for k, v in (
        ("federal minimum wage", 9000), ("federal minimum wage 2026", 5000),
        ("minimum wage usa", 4000), ("minimum wage poster", 3000), ("minimum wage history", 2500))]
    slide = add_programmatic_seo_slide(Presentation(), rows)
    text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame) if slide else ""
    assert "SERP not validated" not in text and "folded in" not in text


def _handshake_rows(url):
    from app.services.keyword_intelligence_service import annotate_keyword_rows
    rows = [
        {"keyword": "handshake deal", "search_volume": 170, "current_position": 7, "current_url": url},
        {"keyword": "handshake agreement", "search_volume": 170, "current_position": 12, "current_url": url},
    ]
    annotate_keyword_rows(rows)
    return rows


def test_ranking_blog_page_is_kept_when_intent_was_only_a_guess():
    # "handshake deal" has no intent marker, so detect_intent guesses
    # Commercial (confidence 55). That guess must not reject the blog post
    # that already ranks #7 for it.
    from app.services.keyword_intelligence_service import ranking_page_target
    from app.services.keyword_relevance_service import classify_page_type
    url = "https://www.lumberfi.com/blog/the-risks-of-handshake-deals-in-construction"
    rows = _handshake_rows(url)
    assert rows[0]["intent_family"] == "Commercial" and rows[0]["intent_confidence"] <= 55
    found = ranking_page_target(rows, None, None, classify_page_type, "Commercial")
    assert found and found["url"] == url and found["match_strength"] == "strong"


def test_confident_commercial_intent_still_rejects_a_blog_post():
    from app.services.keyword_intelligence_service import ranking_page_target
    from app.services.keyword_relevance_service import classify_page_type
    rows = [{"keyword": "buy payroll software", "search_volume": 900, "current_position": 8,
             "current_url": "https://x.com/blog/payroll-tips"}]
    from app.services.keyword_intelligence_service import annotate_keyword_rows
    annotate_keyword_rows(rows)
    assert rows[0]["intent_confidence"] >= 80
    assert ranking_page_target(rows, None, None, classify_page_type, "Commercial") is None


def test_deal_is_a_buying_signal_only_in_a_shopping_phrase():
    from app.services.keyword_intelligence_service import detect_intent
    for shopping in ("laptop deals", "best deal on laptops", "great deal"):
        assert detect_intent(shopping)["intent"] in ("Transactional", "Commercial Investigation")
    for plain in ("handshake deal", "business deal", "deal breaker"):
        assert detect_intent(plain)["intent"] != "Transactional"


def test_branded_example_skips_implausible_zero_ctr_query():
    from app.reporting.pptx_builder import build_branded_dependency_narrative, build_branded_vs_nonbranded_comparison
    nonbranded = [
        {"query": "safety officer", "impressions": 10060, "clicks": 0, "ctr": 0.0, "position": 5.7},
        {"query": "t5018 form", "impressions": 2000, "clicks": 20, "ctr": 0.01, "position": 4.0},
    ]
    comparison = build_branded_vs_nonbranded_comparison(
        [{"query": "lumberfi", "impressions": 900, "clicks": 300, "ctr": 0.33, "position": 1.0}], nonbranded,
    )
    example = build_branded_dependency_narrative(comparison, nonbranded)["concrete_example"]
    assert example and example["query"] == "t5018 form"


def test_branded_example_is_omitted_when_only_implausible_rows_exist():
    from app.reporting.pptx_builder import build_branded_dependency_narrative, build_branded_vs_nonbranded_comparison
    nonbranded = [{"query": "safety officer", "impressions": 10060, "clicks": 0, "ctr": 0.0, "position": 5.7}]
    comparison = build_branded_vs_nonbranded_comparison(
        [{"query": "lumberfi", "impressions": 900, "clicks": 300, "ctr": 0.33, "position": 1.0}], nonbranded,
    )
    assert build_branded_dependency_narrative(comparison, nonbranded)["concrete_example"] is None


def test_target_country_codes_reads_overview_text_and_site_prefixes():
    from app.reporting.pptx_builder import target_country_codes
    assert target_country_codes("Primary USA") == ["usa"]
    assert target_country_codes("Global") is None and target_country_codes(None) is None
    assert target_country_codes("United States and Canada") == ["can", "usa"]
    urls = [f"https://x.com/ca/page-{i}" for i in range(3)]
    assert target_country_codes("Primary USA", urls) == ["can", "usa"]
    assert target_country_codes("Primary USA", urls[:2]) == ["usa"]  # 2 pages is not a market


def _countries():
    return [
        {"country": "usa", "impressions": 244799, "clicks": 1636, "ctr": 0.007, "position": 6.0},
        {"country": "idn", "impressions": 3100, "clicks": 20, "ctr": 0.006, "position": 8.0},
    ]


def test_out_of_market_country_gets_no_page_fix_even_with_query_detail():
    from app.reporting.pptx_builder import build_high_potential_countries
    detail = [{"country": "idn", "page": "https://x.com/blog/books", "query": "pembukuan", "impressions": 900,
               "clicks": 10, "position": 5.0}]
    result = build_high_potential_countries(_countries(), target_countries=["usa"], country_page_query_rows=detail)
    by_country = {r["country"]: r for r in result["material"]}
    assert "outside the client's stated target markets" in by_country["Indonesia"]["fix"]
    assert by_country["Indonesia"]["fix_url"] is None
    assert "outside the client's stated target markets" not in by_country["United States"]["fix"]


def test_country_fix_carries_full_page_url_and_slide_links_it():
    from pptx import Presentation
    from app.reporting.pptx_builder import add_search_opportunities_countries_slide, build_high_potential_countries
    page = "https://www.lumberfi.com/blog/how-to-celebrate-national-construction-appreciation-week-2025-a-guide"
    detail = [{"country": "usa", "page": page, "query": "construction appreciation week", "impressions": 900,
               "clicks": 30, "position": 2.0}]
    result = build_high_potential_countries(_countries()[:1], target_countries=["usa"], country_page_query_rows=detail)
    row = result["material"][0]
    assert row["fix_url"] == page and row["fix_page_label"] in row["fix"]
    slide = add_search_opportunities_countries_slide(Presentation(), result, "Google Search Console")
    links = [
        run.hyperlink.address for sh in slide.shapes if sh.has_table
        for r in sh.table.rows for c in r.cells for p in c.text_frame.paragraphs for run in p.runs
        if run.hyperlink.address
    ]
    assert page in links


def _summary(name, volumes, **extra):
    rows = [{"keyword": f"{name} {i}", "search_volume": v, "detected_intent": "Informational",
             "relevance_status": "Core Relevant", "keyword_difficulty": 10, "current_position": 5,
             "domain_positions": {"rival.com": 3}} for i, v in enumerate(volumes)]
    return {"name": name, "rows": rows, "confidence_level": "High", **extra}


def test_tiny_one_keyword_cluster_is_not_high_priority():
    from app.services.keyword_strategy_service import score_summaries
    summaries = [_summary("Big A", [40000, 9000]), _summary("Big B", [30000, 5000]), _summary("Big C", [20000]),
                 _summary("Handshake", [170])]
    score_summaries(summaries)
    tiers = {s["name"]: s["roadmap_priority"] for s in summaries}
    assert tiers["Handshake"] != "High"
    assert "High" in tiers.values()


def test_small_market_client_can_still_have_high_priority_topics():
    from app.services.keyword_strategy_service import score_summaries
    summaries = [_summary("Topic A", [90, 40]), _summary("Topic B", [60]), _summary("Topic C", [50])]
    score_summaries(summaries)
    assert any(s["roadmap_priority"] == "High" for s in summaries)

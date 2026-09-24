"""Target Keywords layout (2026-09-24 spec): "Target Keywords" heading,
"(Category)" only on a category's first slide, a cluster subheading per
table, columns exactly Clusters | Keywords | Volume | KD | Intent, and a
real vertically merged Clusters cell over each cluster's keyword rows."""

from pptx import Presentation

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, TARGET_KEYWORDS_HEADERS, _audit_slide_geometry, _audit_target_keyword_slides,
    _render_target_keyword_slides, add_keyword_research_slide, add_strategic_keyword_clusters_slide,
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _named(slide, name):
    return [sh.text_frame.text for sh in slide.shapes if sh.name == name]


def _tables(slide):
    return [sh.table for sh in slide.shapes if sh.has_table]


def _cluster(name, n, prefix=None):
    return {
        "name": name, "source": "test", "insights": [f"{name} insight"],
        "rows": [(f"{prefix or name.lower()} keyword {i}", f"{1000 - i:,}", str(20 + i), "Commercial") for i in range(n)],
    }


def test_heading_category_subheading_and_merged_cell():
    sheet = [{"cluster": "Trucks", "keywords": [
        {"keyword": "heavy truck price", "search_volume": 110000, "keyword_difficulty": 30, "intent": "Commercial", "sub_category": "Heavy Duty Trucks"},
        {"keyword": "heavy duty truck", "search_volume": 14800, "keyword_difficulty": 28, "intent": "Commercial", "sub_category": "Heavy Duty Trucks"},
        {"keyword": "what is a heavy truck", "search_volume": 4400, "keyword_difficulty": 25, "intent": "Informational", "sub_category": "Heavy Duty Trucks"},
        {"keyword": "10 wheeler truck", "search_volume": 3600, "keyword_difficulty": 22, "intent": "Commercial", "sub_category": "Heavy Duty Trucks"},
        {"keyword": "tipper truck", "search_volume": 9000, "keyword_difficulty": 35, "intent": "Commercial", "sub_category": "Truck Types & Applications"},
        {"keyword": "tractor trailer", "search_volume": 5000, "keyword_difficulty": 31, "intent": "Commercial", "sub_category": "Truck Types & Applications"},
    ]}]
    slides = add_strategic_keyword_clusters_slide(_prs(), sheet)
    slide = slides[0]
    first_text = next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    assert first_text == "Target Keywords"
    assert _named(slide, "TK Category") == ["(Trucks)"]
    subheads = [n for s in slides for n in _named(s, "TK Cluster")]
    assert subheads == ["Heavy Duty Trucks", "Truck Types & Applications"]
    table = _tables(slide)[0]
    assert [table.cell(0, j).text for j in range(len(table.columns))] == TARGET_KEYWORDS_HEADERS
    origin = table.cell(1, 0)
    assert origin.text == "Heavy Duty Trucks"
    assert origin.is_merge_origin and origin.span_height == 4
    assert all(table.cell(i, 0).is_spanned for i in range(2, 5))
    assert [table.cell(i, 1).text for i in range(1, 5)] == [
        "heavy truck price", "heavy duty truck", "what is a heavy truck", "10 wheeler truck",
    ]
    assert table.cell(1, 2).text == "110,000" and table.cell(3, 4).text == "Informational"


def test_category_line_only_on_first_slide_of_each_category():
    categories = [
        {"name": "Trucks", "clusters": [_cluster("Heavy Duty Trucks", 9), _cluster("Tippers", 8)]},
        {"name": "Buses", "clusters": [_cluster("School Buses", 9)]},
    ]
    slides = _render_target_keyword_slides(_prs(), categories)
    assert [_named(s, "TK Category") for s in slides] == [["(Trucks)"], [], ["(Buses)"]]
    assert _audit_target_keyword_slides(slides, categories) == []


def test_long_cluster_continues_without_losing_rows():
    categories = [{"name": "Trucks", "clusters": [_cluster("Heavy Duty Trucks", 30)]}]
    prs = _prs()
    slides = _render_target_keyword_slides(prs, categories)
    assert len(slides) > 1
    assert all(_named(s, "TK Cluster") == ["Heavy Duty Trucks"] for s in slides)
    assert sum(len(t.rows) - 1 for s in slides for t in _tables(s)) == 30
    assert _audit_target_keyword_slides(slides, categories) == []
    assert _audit_slide_geometry(prs) == []


def test_two_small_clusters_share_a_slide_with_separate_merged_cells():
    categories = [{"name": "Trucks", "clusters": [_cluster("Tippers", 3), _cluster("Tractors", 2)]}]
    slides = _render_target_keyword_slides(_prs(), categories)
    assert len(slides) == 1
    tables = _tables(slides[0])
    assert [t.cell(1, 0).text for t in tables] == ["Tippers", "Tractors"]
    assert tables[0].cell(1, 0).span_height == 3 and tables[1].cell(1, 0).span_height == 2


def test_small_clusters_never_paired_across_categories_or_without_one():
    categories = [
        {"name": "Trucks", "clusters": [_cluster("Tippers", 2)]},
        {"name": "Buses", "clusters": [_cluster("School Buses", 2)]},
        {"name": None, "clusters": [_cluster("A", 2), _cluster("B", 2)]},
    ]
    slides = _render_target_keyword_slides(_prs(), categories)
    assert [len(_tables(s)) for s in slides] == [1, 1, 1, 1]
    assert _audit_target_keyword_slides(slides, categories) == []


def test_audit_catches_changed_or_missing_rows():
    categories = [{"name": "Trucks", "clusters": [_cluster("Tippers", 3)]}]
    slides = _render_target_keyword_slides(_prs(), categories)
    tampered = [{"name": "Trucks", "clusters": [_cluster("Tippers", 4)]}]
    assert any("differ from source" in p for p in _audit_target_keyword_slides(slides, tampered))


def test_ai_path_uses_business_theme_as_category_and_new_columns():
    rows = [
        {"keyword": "payroll software", "cluster": "Payroll Software", "search_volume": 900, "keyword_difficulty": 40,
         "detected_intent": "Commercial", "business_theme": "Payroll"},
        {"keyword": "payroll app", "cluster": "Payroll Software", "search_volume": 300, "keyword_difficulty": 30,
         "detected_intent": "Commercial", "business_theme": "Payroll"},
    ]
    prs = _prs()
    slides = add_keyword_research_slide(prs, rows)
    assert _named(slides[0], "TK Category") == ["(Payroll)"]
    table = _tables(slides[0])[0]
    assert [table.cell(0, j).text for j in range(5)] == TARGET_KEYWORDS_HEADERS
    assert table.cell(1, 0).text == "Payroll Software" and table.cell(1, 0).span_height == 2
    assert _audit_slide_geometry(prs) == []


def test_worst_case_long_names_never_overflow():
    long_name = "Very Long Descriptive Cluster Topic Name That Wraps Inside The Merged Cell"
    categories = [{"name": "Category With A Fairly Long Name", "clusters": [
        {**_cluster(long_name, 1), "rows": [("a very long descriptive keyword phrase " * 3, "1", "2", "Commercial")]},
        _cluster("Second", 14),
    ]}]
    prs = _prs()
    slides = _render_target_keyword_slides(prs, categories)
    assert _audit_target_keyword_slides(slides, categories) == []
    assert _audit_slide_geometry(prs) == []

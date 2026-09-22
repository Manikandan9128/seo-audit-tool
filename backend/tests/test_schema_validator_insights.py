from pptx import Presentation

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, add_schema_combined_slide, build_schema_report_parts, schema_eligibility_notes,
)
from app.services.technical_seo_service import aggregate_schema_validation


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def _pages(blog_n=310, product_n=55, other_n=1596):
    pages = []
    for i in range(blog_n):
        pages.append({"url": f"https://x.com/blog/post-{i}", "meta": {"schema_types_found": [], "schema_field_issues": []}})
    for i in range(product_n):
        pages.append({"url": f"https://x.com/product/item-{i}", "meta": {"schema_types_found": [], "schema_field_issues": []}})
    for i in range(other_n):
        pages.append({"url": f"https://x.com/page-{i}", "meta": {"schema_types_found": ["WebSite", "Organization"], "schema_field_issues": []}})
    return pages


def test_other_pages_marked_na_not_a_content_schema_row():
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    other = next(r for r in parts["part1"] if r["page_type"] == "Other Pages")
    assert other["recommended_schema"] == "N/A"
    assert not any(r["schema_type"] == "Other Pages" for r in parts["part2"])


def test_part1_order_is_fixed_regardless_of_upstream_pageview_sort():
    # 2026-09-22 spec rule 1: Blog/Article, Product, JobPosting, other
    # types, Site-wide, Other Pages always last — regardless of what order
    # aggregate_schema_validation's own by_page_type arrives in (it's
    # pageview/priority-sorted for its OTHER consumers, and here "Other
    # Pages" is deliberately listed FIRST to simulate the exact case that
    # sort could produce if Other Pages happened to have the most traffic).
    schema_validation = {
        "total_pages": 500,
        "by_page_type": [
            {"page_type": "Other Pages", "pages": 300, "present_pages": 0, "valid_pages": 0, "valid_pct": 0},
            {"page_type": "LocalBusiness", "pages": 10, "present_pages": 0, "valid_pages": 0, "valid_pct": 0},
            {"page_type": "JobPosting", "pages": 5, "present_pages": 0, "valid_pages": 0, "valid_pct": 0},
            {"page_type": "Product", "pages": 80, "present_pages": 0, "valid_pages": 0, "valid_pct": 0},
            {"page_type": "Article-type", "pages": 105, "present_pages": 0, "valid_pages": 0, "valid_pct": 0},
        ],
        "type_coverage": [], "missing_properties": [],
    }
    parts = build_schema_report_parts(schema_validation)
    page_types = [r["page_type"] for r in parts["part1"]]
    assert page_types == ["Blog / Article", "Product", "Job Detail Pages", "Local / Location", "Site-wide", "Other Pages"]


def test_jobposting_bucket_excludes_careers_index_department_and_recruitment_pages():
    # 2026-09-22 spec rule 2: only individual job-detail pages qualify —
    # careers landing/index, department/team, and general recruitment
    # pages must never be classified as JobPosting.
    pages = [
        {"url": "https://x.com/careers/", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/openings", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/engineering-department", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/join-our-team", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/recruitment-process", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/senior-backend-engineer-4821", "meta": {"schema_types_found": [], "schema_field_issues": []}},
    ]
    sv = aggregate_schema_validation(pages)
    job_rows = [r for r in sv["by_page_type"] if r["page_type"] == "JobPosting"]
    assert len(job_rows) == 1
    assert job_rows[0]["pages"] == 1  # only the real job-detail page


def test_jobposting_why_it_applies_is_never_blank():
    schema_validation = {
        "total_pages": 10,
        "by_page_type": [{"page_type": "JobPosting", "pages": 3, "present_pages": 0, "valid_pages": 0, "valid_pct": 0}],
        "type_coverage": [], "missing_properties": [],
    }
    parts = build_schema_report_parts(schema_validation)
    job_row = next(r for r in parts["part1"] if r["page_type"] == "Job Detail Pages")
    assert job_row["why_it_applies"]
    assert "careers index/listing page" in job_row["why_it_applies"]


def test_content_type_coverage_denominator_is_its_own_page_count_not_total():
    # Part 2's Coverage % denominator must be the pages that type applies
    # to (310 blog pages), never the full site total (~1961 pages).
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    article_row = next(r for r in parts["part2"] if r["schema_type"] == "Article")
    assert article_row["applicable"] == 310
    product_row = next(r for r in parts["part2"] if r["schema_type"] == "Product")
    assert product_row["applicable"] == 55


def test_baseline_schema_is_site_level_not_a_page_count():
    # 2026-09-20 spec: WebSite/Organization are entity-level facts, never
    # counted per crawled page.
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    website_row = next(r for r in parts["part2"] if r["schema_type"] == "WebSite")
    assert website_row["applicable"] == "Site-level"
    assert website_row["present"] == "Yes"
    assert website_row["missing"] == "—"
    assert website_row["coverage_pct"] is None
    org_row = next(r for r in parts["part2"] if r["schema_type"] == "Organization")
    assert org_row["applicable"] == "Site-level"
    assert org_row["present"] == "Yes"


def test_breadcrumb_applicable_is_content_pages_only_not_other_pages():
    # BreadcrumbList's denominator should be real content pages (blog +
    # product), matching the SNIP reference — not blended with the "Other
    # Pages" utility bucket, and not the full site total either.
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    breadcrumb_row = next(r for r in parts["part2"] if r["schema_type"] == "BreadcrumbList")
    assert breadcrumb_row["applicable"] == 310 + 55


def test_site_wide_part1_row_present():
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    site_wide = next(r for r in parts["part1"] if r["page_type"] == "Site-wide")
    assert site_wide["recommended_schema"] == "WebSite, Organization"
    # 2026-09-20 spec: site-level Pages reads "Site-level", never a page count.
    assert site_wide["pages"] == "Site-level"


def test_jobposting_never_appears_with_no_job_detail_pages():
    sv = aggregate_schema_validation(_pages())
    parts = build_schema_report_parts(sv)
    assert not any("JobPosting" in str(v) for r in parts["part1"] for v in r.values())
    assert not any("JobPosting" in str(v) for r in parts["part2"] for v in r.values())


def test_jobposting_never_appears_for_careers_index_page_alone():
    # Universal SEO Audit Engine spec (2026-09-20) section 29: a careers
    # index/listing page is NOT itself a JobPosting page.
    pages = _pages() + [
        {"url": "https://x.com/careers/", "meta": {"schema_types_found": [], "schema_field_issues": []}},
        {"url": "https://x.com/careers/openings", "meta": {"schema_types_found": [], "schema_field_issues": []}},
    ]
    sv = aggregate_schema_validation(pages)
    parts = build_schema_report_parts(sv)
    assert not any("JobPosting" in str(v) for r in parts["part2"] for v in r.values())


def test_jobposting_appears_for_real_job_detail_page():
    pages = _pages() + [
        {"url": "https://x.com/careers/senior-backend-engineer", "meta": {"schema_types_found": ["JobPosting"], "schema_field_issues": []}},
    ]
    sv = aggregate_schema_validation(pages)
    parts = build_schema_report_parts(sv)
    job_row = next(r for r in parts["part2"] if r["schema_type"] == "JobPosting")
    assert job_row["applicable"] == 1
    assert job_row["present"] == 1


def test_content_type_reports_present_valid_invalid_missing_separately():
    pages = [
        {"url": "https://x.com/blog/a", "meta": {"schema_types_found": ["Article"], "schema_field_issues": []}},
        {"url": "https://x.com/blog/b", "meta": {"schema_types_found": ["Article"], "schema_field_issues": ["Article schema missing required field: image"]}},
        {"url": "https://x.com/blog/c", "meta": {"schema_types_found": [], "schema_field_issues": []}},
    ]
    sv = aggregate_schema_validation(pages)
    parts = build_schema_report_parts(sv)
    article_row = next(r for r in parts["part2"] if r["schema_type"] == "Article")
    assert article_row["applicable"] == 3
    assert article_row["present"] == 2
    assert article_row["valid"] == 1
    assert article_row["invalid"] == 1
    assert article_row["missing"] == 1


def test_eligibility_notes_flag_retired_type_only():
    sv = aggregate_schema_validation(_pages(blog_n=0, product_n=0, other_n=50))
    parts = build_schema_report_parts(sv)
    notes = schema_eligibility_notes(parts["part2"])
    assert "Product" not in notes
    assert "WebSite" not in notes


def test_slide_renders_with_no_ai_insights_and_no_part3_section():
    sv = aggregate_schema_validation(_pages())
    slide = add_schema_combined_slide(_prs(), sv, schema_ai_insights=None)
    text = _slide_text(slide)
    assert "Part 1" in text
    assert "Part 2" in text
    assert "KEY INSIGHTS" not in text.upper()


def test_slide_renders_ai_insights_when_provided():
    sv = aggregate_schema_validation(_pages())
    slide = add_schema_combined_slide(_prs(), sv, schema_ai_insights={"insights": ["Article schema missing on all 310 blog pages — Fix: add it to the blog template."]})
    text = _slide_text(slide)
    assert "Fix: add it to the blog template" in text

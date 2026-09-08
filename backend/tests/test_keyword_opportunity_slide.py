from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _ctr_for_position, add_keyword_opportunity_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_ctr_benchmark_known_positions():
    assert _ctr_for_position(1) == 0.317
    assert _ctr_for_position(10) == 0.022
    assert _ctr_for_position(25) == 0.01  # beyond page 1
    assert _ctr_for_position(0) == 0.0
    assert _ctr_for_position(None) == 0.0


def test_returns_none_when_no_clusters():
    assert add_keyword_opportunity_slide(_prs(), [{"keyword": "x", "search_volume": 100}]) is None


def test_returns_none_when_all_clusters_zero_volume():
    rows = [{"keyword": "x", "cluster": "A", "search_volume": 0}]
    assert add_keyword_opportunity_slide(_prs(), rows) is None


def test_creates_slide_for_valid_cluster():
    rows = [{"keyword": "certified payroll software", "cluster": "Certified Payroll", "search_volume": 1600, "keyword_difficulty": 35}]
    assert add_keyword_opportunity_slide(_prs(), rows) is not None


def test_not_ranking_no_page_recommends_create_new_page():
    rows = [{"keyword": "certified payroll software", "cluster": "A", "search_volume": 1600, "keyword_difficulty": 35}]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    assert slide is not None
    text = _slide_text(slide)
    assert "Create New Page" in text


def test_not_ranking_with_existing_page_recommends_realign():
    rows = [{
        "keyword": "construction payroll compliance", "cluster": "A", "search_volume": 800,
        "keyword_difficulty": 40, "current_url": "https://example.com/compliance",
    }]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "Re-Align Page" in text


def test_ranking_beyond_page_one_recommends_optimize():
    rows = [{"keyword": "labor burden calculator", "cluster": "A", "search_volume": 400, "keyword_difficulty": 25, "current_position": 15}]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "Optimize Existing Page" in text


def test_ranking_page_one_not_top_three_recommends_expand():
    rows = [{"keyword": "union payroll software", "cluster": "A", "search_volume": 300, "keyword_difficulty": 30, "current_position": 6}]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "Expand Content" in text


def test_ranking_top_three_recommends_monitor():
    rows = [{"keyword": "prevailing wage software", "cluster": "A", "search_volume": 500, "keyword_difficulty": 20, "current_position": 2}]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "Monitor" in text


def test_priority_tiers_split_across_multiple_clusters():
    rows = [
        {"keyword": f"keyword {i}", "cluster": f"Cluster {i}", "search_volume": 1000 - i * 10, "keyword_difficulty": 30}
        for i in range(9)
    ]
    slide = add_keyword_opportunity_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "High" in text
    assert "Medium" in text
    assert "Low" in text

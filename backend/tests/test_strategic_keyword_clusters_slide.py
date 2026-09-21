from pptx import Presentation

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, _audit_slide_geometry, add_strategic_keyword_clusters_slide,
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_returns_empty_list_when_no_data():
    assert add_strategic_keyword_clusters_slide(_prs(), None) == []
    assert add_strategic_keyword_clusters_slide(_prs(), []) == []


def test_one_slide_per_cluster_with_real_values():
    clusters = [
        {
            "cluster": "Heavy Trucks",
            "keywords": [
                {"keyword": "6x4 truck price", "search_volume": 1200, "keyword_difficulty": 40, "intent": "Commercial", "sub_category": "Pricing"},
                {"keyword": "6x4 truck specs", "search_volume": 800, "keyword_difficulty": 30, "sub_category": "Specifications"},
            ],
        },
        {
            "cluster": "Fleet Financing",
            "keywords": [{"keyword": "truck loan emi", "search_volume": 500, "keyword_difficulty": 25}],
        },
    ]
    slides = add_strategic_keyword_clusters_slide(_prs(), clusters)
    assert len(slides) == 2
    text0 = _slide_text(slides[0])
    assert "Heavy Trucks" in text0
    assert "6x4 truck price" in text0
    assert "1,200" in text0
    assert "40" in text0
    assert "Commercial" in text0
    assert "Pricing" in text0


def test_missing_optional_fields_render_as_dash_not_zero():
    clusters = [{
        "cluster": "Widgets",
        "keywords": [
            {"keyword": "widget pricing"},  # no volume/kd/intent/sub_category at all
        ],
    }]
    slides = add_strategic_keyword_clusters_slide(_prs(), clusters)
    text = _slide_text(slides[0])
    assert "widget pricing" in text
    assert "—" in text  # em-dash placeholder, never a fabricated 0


def test_no_overlap_worst_case_six_clusters_eight_keywords_each():
    clusters = []
    for c in range(6):
        clusters.append({
            "cluster": f"Very Long Descriptive Cluster Topic Number {c}",
            "keywords": [
                {
                    "keyword": f"very long descriptive keyword phrase number {c} {k}",
                    "search_volume": 50000 - c * 100 - k, "keyword_difficulty": 40,
                    "intent": "Commercial", "sub_category": f"Sub Topic {k % 3}",
                }
                for k in range(8)
            ],
        })
    prs = _prs()
    add_strategic_keyword_clusters_slide(prs, clusters)
    assert _audit_slide_geometry(prs) == []

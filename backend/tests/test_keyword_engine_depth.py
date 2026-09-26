"""Universal SEO Keyword prompt — 2026-09-24 depth pass (keyword_site_model +
keyword_strategy_depth): the website model and everything it feeds (§3, §4,
§6, §11, §12, §18, §19/§20, §23, §32, §36-§38, §46-§52, §54-§59, §62, §66).
See docs/keyword_engine_spec_coverage.md."""

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _audit_slide_geometry, add_keyword_master_slide
from app.services.google_sheets_service import keyword_strategy_tabs
from app.services.keyword_intelligence_service import annotate_keyword_rows
from app.services.keyword_site_model import _target_page_authority, build_site_model, geo_served, keyword_audience_from_site
from app.services.keyword_semantic_signals import build_corpus_idf
from app.services.keyword_strategy_depth import (
    cannibalization_similarity, content_gaps, deepen_topics, entity_relationship, extra_review_items,
    programmatic_patterns,
)
from app.services.keyword_strategy_service import build_full_keyword_strategy
from pptx import Presentation


def _pages():
    return [
        {"page_url": "https://acme.com/", "page_title": "Acme HR Platform", "word_count": 400},
        {"page_url": "https://acme.com/services/payroll", "page_title": "Payroll Processing Services",
         "word_count": 1800, "incoming_internal_links": 12, "crawl_depth": 1},
        {"page_url": "https://acme.com/services/payroll-2", "page_title": "Payroll Processing Service",
         "word_count": 1700, "incoming_internal_links": 3, "crawl_depth": 2},
        {"page_url": "https://acme.com/industries/construction", "page_title": "Payroll for Construction",
         "word_count": 900, "incoming_internal_links": 4, "crawl_depth": 2},
        {"page_url": "https://acme.com/locations/austin-tx", "page_title": "Payroll Services in Austin",
         "word_count": 500, "incoming_internal_links": 2, "crawl_depth": 2},
        {"page_url": "https://acme.com/blog/employee-compensation", "page_title": "Employee Compensation Glossary",
         "word_count": 1200, "incoming_internal_links": 6, "crawl_depth": 2},
    ]


def _overview():
    return {"company_name": "Acme", "solutions": ["payroll", "benefits"], "industries": ["construction"],
            "primary_buyers": ["small business owners"], "target_country": "United States"}


def _summary(name, kws, family="Commercial", url=None, strength="none", **extra):
    rows = [{"keyword": k, "search_volume": v, "keyword_difficulty": 30} for k, v in kws]
    annotate_keyword_rows(rows)
    for r in rows:
        r.pop("_core_key", None), r.pop("_core_tokens", None)
        r["intent_family"] = r.get("intent_family") if family is None else family
    return {"name": name, "rows": rows, "target_url": url, "match_strength": strength,
            "confidence_level": "High", "dominant_family": family, "primary_keyword": kws[0][0], **extra}


# §3/§4 website model ----------------------------------------------------------------
def test_site_model_builds_entity_graph_and_14_url_fields():
    model = build_site_model(_pages(), _overview(), {"https://acme.com/services/payroll": 400})
    assert "payroll" in model["graph"]["services"]
    assert "construction" in model["graph"]["industries"]
    page = next(p for p in model["pages"] if p["url"] == "https://acme.com/services/payroll")
    for field in ("page_type", "primary_entity", "topic", "audience", "intent", "funnel_stage", "geography",
                  "content_depth", "incoming_internal_links", "crawl_depth", "internal_link_role", "business_purpose",
                  "traffic_clicks"):
        assert field in page, field
    assert page["content_depth"] == "Deep" and page["traffic_clicks"] == 400


# §23 per-page authority, grouped from the same Backlinks export already parsed --------
def test_page_authority_from_backlink_rows():
    rows = [
        {"target_url": "https://acme.com/services/payroll", "source_url": "https://a.com/post", "domain_score": 60},
        {"target_url": "https://acme.com/services/payroll", "source_url": "https://b.com/post", "domain_score": 40},
        {"target_url": "https://acme.com/services/payroll/", "source_url": "https://a.com/other-post", "domain_score": 60},
        {"target_url": "https://acme.com/blog/x", "source_url": "https://c.com/post", "domain_score": 20},
    ]
    authority = _target_page_authority(rows)
    payroll = authority["https://acme.com/services/payroll"]
    assert payroll["referring_backlinks"] == 2  # a.com counted once despite two links
    assert payroll["authority_score"] == 53  # round((60+40+60)/3)

    model = build_site_model(_pages(), _overview(), backlink_rows=rows)
    page = next(p for p in model["pages"] if p["url"] == "https://acme.com/services/payroll")
    assert page["authority_score"] == 53 and page["referring_backlinks"] == 2


# §24/§58 duplicate-title cannibalization from the crawl alone ----------------------
def test_duplicate_titles_flag_near_identical_pages():
    model = build_site_model(_pages(), _overview())
    dupes = model["duplicate_titles"]
    assert any(d["preferred_url"] == "https://acme.com/services/payroll" and
               d["other_url"] == "https://acme.com/services/payroll-2" for d in dupes)


# §11/§12 site-aware audience + geography -------------------------------------------
def test_audience_and_geo_read_from_the_sites_own_pages():
    model = build_site_model(_pages(), _overview())
    assert keyword_audience_from_site("payroll software for construction crews", model) == "Construction"
    assert geo_served("austin", model) == "Served (listed location)"
    assert geo_served("united states", model) == "Served (country-level)"
    assert geo_served("tokyo", model) == "Not listed — check before targeting"
    assert geo_served(None, model) is None


# §6 entity relationship -------------------------------------------------------------
def test_entity_relationship_labels():
    assert entity_relationship({"keyword": "payroll vs gusto", "detected_intent": "Comparison"}, {}) == "product -> alternative"
    assert entity_relationship({"keyword": "payroll for construction", "entity_type": "Service"}, {"industries": ["construction"]}) \
        == "service -> industry"


# End-to-end: full strategy with a real site model wired in -------------------------
def test_full_strategy_with_site_model():
    summaries = [
        _summary("Payroll Services", [("payroll services", 2000), ("payroll service", 1500)],
                 url="https://acme.com/services/payroll", strength="strong"),
        _summary("Payroll For Construction", [("payroll for construction companies", 400)], url=None, strength="none"),
        _summary("What Is Payroll", [("what is payroll", 900)], family="Informational"),
    ]
    backlinks = [{"target_url": "https://acme.com/services/payroll", "source_url": "https://a.com", "domain_score": 55}]
    ctx = {
        "site_authority": 40, "brands": {"own": {"acme"}},
        "site_audit_pages_rows": _pages(), "company_overview": _overview(),
        "page_clicks": {"https://acme.com/services/payroll": 400},
        "site_model": build_site_model(_pages(), _overview(), {"https://acme.com/services/payroll": 400}, backlinks),
        "page_query_rows": [], "routed_counts": {},
    }
    strategy = build_full_keyword_strategy(summaries, ctx)
    assert strategy["site"]["graph"] and strategy["site"]["pages"]

    row = summaries[0]["rows"][0]
    for field in ("audience", "geo_served", "entity_relationship", "language", "existing_url_traffic",
                  "topical_authority", "primary_score", "primary_or_secondary"):
        assert field in row, field
    assert row["language"] == "English / Latin script"
    assert 0.0 <= row["topical_authority"] <= 1.0
    assert summaries[0]["primary_keyword"] in {"payroll services", "payroll service"}

    cluster = strategy["clusters"][0]
    for field in ("audience", "geography", "difficulty", "search_demand", "business_value", "target_evidence",
                  "secondary_keywords", "keyword_count"):
        assert field in cluster, field
    assert cluster["target_evidence"]["authority_score"] == 55

    page = next(p for p in strategy["page_map"] if p["url"] == "https://acme.com/services/payroll")
    assert page["primary_entity"] and "supporting_topics" in page

    tabs = dict(keyword_strategy_tabs(strategy))
    assert {"Clusters", "Site Pages", "Entity Graph"} <= set(tabs)
    assert tabs["Clusters"][0][0] == "Cluster ID"


# §47 local-value business rule -------------------------------------------------------
def test_local_page_without_service_areas_goes_to_review():
    s = _summary("Payroll In Miami", [("payroll services miami", 300)], family="Local")
    s["decision"], s["decision_reason"] = "NEW URL REQUIRED", "no match"
    from app.services.keyword_strategy_depth import apply_business_rules
    apply_business_rules([s], {"service_areas": [], "pages": []}, {"models": ["Service"]}, set())
    assert s["decision"] == "REVIEW" and s["business_rule"] == "§47 local value check"


# §57 content gap output has all the required fields ---------------------------------
def test_content_gap_output_fields():
    s = _summary("Payroll Comparisons", [("best payroll software", 800), ("payroll software comparison", 300)],
                  family="Comparison")
    s["content_gap_type"] = "Comparison Content Gap"
    s["page_type"] = "Comparison Page"
    s["roadmap_priority"] = "High"
    for r in s["rows"]:
        r["domain_positions"] = {"competitor.com": 4}
    gaps = content_gaps([s], {"sections": {"comparison": "/compare/"}})
    assert gaps and gaps[0]["suggested_url"].startswith("/compare/")
    for field in ("business_relevance", "competitor_evidence", "search_demand", "supporting_keywords", "reason"):
        assert field in gaps[0], field


# §58 cannibalization similarity + duplicate-title pairs ------------------------------
def test_cannibalization_similarity_and_duplicate_pairs():
    canni = [{"query": "q", "preferred_url": "https://acme.com/services/payroll",
              "other_urls": ["https://acme.com/services/payroll-2"], "risk": "High", "action": "Merge",
              "evidence": "e", "impressions": 100}]
    model = build_site_model(_pages(), _overview())
    out = cannibalization_similarity(canni, model)
    assert out[0]["similarity"] > 0 and out[0]["source"] == "Search Console"
    assert any(c["source"] == "Duplicate titles" for c in out)


def test_cannibalization_merges_multiple_duplicates_of_the_same_preferred_page():
    # Regression (confirmed real, Geopits regen 2026-09-26): a paginated
    # blog series (?page=1 preferred, ?page=2 and ?page=3 near-identical
    # siblings) used to emit one dict per PAIR, both sharing the same
    # preferred_url — since the old evidence text never named either page
    # ("Two blog pages with 92% title overlap."), the two pairs rendered
    # as byte-identical bullets on the Next Steps: Content SEO slide.
    # Siblings of the same preferred page must merge into ONE finding,
    # with every sibling actually named.
    fake_site_model = {
        "duplicate_titles": [
            {"preferred_url": "https://x.com/blog?page=1", "other_url": "https://x.com/blog?page=2",
             "similarity": 0.92, "page_type": "blog"},
            {"preferred_url": "https://x.com/blog?page=1", "other_url": "https://x.com/blog?page=3",
             "similarity": 0.88, "page_type": "blog"},
        ],
    }
    out = cannibalization_similarity([], fake_site_model)
    dupe_findings = [c for c in out if c["source"] == "Duplicate titles"]
    assert len(dupe_findings) == 1, f"expected one merged finding, got {len(dupe_findings)}: {dupe_findings}"
    finding = dupe_findings[0]
    assert set(finding["other_urls"]) == {"https://x.com/blog?page=2", "https://x.com/blog?page=3"}
    assert "https://x.com/blog?page=2" in finding["evidence"]
    assert "https://x.com/blog?page=3" in finding["evidence"]


# §27/§59/§62 programmatic patterns + extra review types ------------------------------
def test_programmatic_patterns_and_review_items():
    s = _summary("Payroll By Business Type", [
        ("payroll for small business", 400), ("payroll for enterprise", 300), ("payroll for contractors", 250),
        ("payroll for business", 200)], family="Commercial")
    s["cluster_source"] = "rule"
    patterns = programmatic_patterns([s])
    assert patterns and patterns[0]["dimension_b"] == "Audience"
    assert patterns[0]["serp_validation"] == "Not available (no SERP data source yet)"
    items = extra_review_items([s], patterns)
    assert any(i["type"] == "Programmatic opportunity" for i in items)


# §66-B fallback slide when Google Sheets isn't connected ------------------------------
def test_keyword_master_fallback_slide_when_no_sheet():
    strategy = {"clusters": [{"cluster_id": "C01", "name": "Payroll Services", "primary_keyword": "payroll services",
                              "intent": "Commercial", "priority": "High", "decision": "existing url", "parent_topic": "Payroll"}],
                "topics": []}
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = add_keyword_master_slide(prs, strategy)
    assert slide is not None
    assert _audit_slide_geometry(prs) == []
    assert add_keyword_master_slide(prs, None) is None


def test_keyword_master_shows_as_many_rows_as_it_claims():
    # Regression (confirmed real, Geopits regen 2026-09-26): the slide's
    # own insights line said "the top 12 by priority are shown", but
    # _draw_table's row_cap defaults to 9 (not 14) whenever `insights` is
    # passed — a generic heuristic never checked against this table's own
    # columns — so only 9 of the 12 computed rows actually rendered.
    clusters = [
        {"cluster_id": f"C{i:02d}", "name": f"Cloud Database Migration Strategy Cluster {i}",
         "parent_topic": "Cloud & Database Services",
         "primary_keyword": f"managed postgresql database migration services {i}",
         "intent": "Commercial", "priority": "High", "decision": "new_url_required"}
        for i in range(12)
    ]
    strategy = {"clusters": clusters, "topics": []}
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = add_keyword_master_slide(prs, strategy, max_rows=12)
    assert _audit_slide_geometry(prs) == []
    table = next(sh.table for sh in slide.shapes if sh.has_table)
    assert len(table.rows) - 1 == 12, f"claimed top 12 but rendered {len(table.rows) - 1} rows"


# §12/§15 zero-API same-need review signal (SERP_API_Alternatives doc, 2026-09-24) -----
def test_deepen_topics_flags_a_genuine_zero_overlap_same_need_pair():
    a = _summary("Remote DBA", [("remote dba services", 900)], url=None, strength="none")
    b = _summary("Database Support", [("database support services", 700)], url=None, strength="none")
    for s in (a, b):
        for r in s["rows"]:
            r["entity_type"] = "Service"
            r["user_need"] = "hire a database administrator"
    topic = {"parent": "database", "clusters": ["Remote DBA", "Database Support"], "covered": 0, "total": 2, "links": []}
    idf = build_corpus_idf([r["keyword"] for s in (a, b) for r in s["rows"]])
    deepen_topics([topic], [a, b], idf)
    assert topic["semantic_similar_pairs"]
    assert {topic["semantic_similar_pairs"][0]["a"], topic["semantic_similar_pairs"][0]["b"]} == {"Remote DBA", "Database Support"}
    items = extra_review_items([a, b], [], [topic])
    assert any(i["type"] == "Possible same-need clusters (zero-API signal)" for i in items)


def test_deepen_topics_does_not_flag_unrelated_clusters_in_the_same_topic():
    a = _summary("Payroll Pricing", [("payroll pricing", 900)], family="Transactional")
    b = _summary("Payroll Careers", [("payroll company careers", 200)], family="Navigational")
    for s in (a, b):
        for r in s["rows"]:
            r["entity_type"] = None
            r["user_need"] = None
    topic = {"parent": "payroll", "clusters": ["Payroll Pricing", "Payroll Careers"], "covered": 0, "total": 2, "links": []}
    idf = build_corpus_idf([r["keyword"] for s in (a, b) for r in s["rows"]])
    deepen_topics([topic], [a, b], idf)
    assert topic["semantic_similar_pairs"] == []


def test_deepen_topics_is_a_noop_without_idf():
    a = _summary("Remote DBA", [("remote dba services", 900)])
    topic = {"parent": "database", "clusters": ["Remote DBA"], "covered": 0, "total": 1, "links": []}
    deepen_topics([topic], [a])  # no idf passed
    assert topic["semantic_similar_pairs"] == []

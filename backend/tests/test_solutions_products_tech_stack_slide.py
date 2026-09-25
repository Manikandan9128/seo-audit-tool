"""Solutions, Products & Tech Stack slide (spec 2026-09-25): the old
standalone "Tech Stack & Hosting" slide (add_tech_stack_slide) was
removed — its content now lives in this slide's right card. Left card
(Solutions/Products) keeps its previous content, just narrower."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, _audit_slide_geometry, add_solutions_products_slide,
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _texts(slide) -> list[str]:
    return [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def test_none_when_everything_is_empty():
    assert add_solutions_products_slide(_prs(), {}, {}) is None
    assert add_solutions_products_slide(_prs(), {}, None) is None


def test_renders_from_tech_stack_alone_when_overview_is_empty():
    # The old standalone slide only needed tech_stack — folding it into
    # this slide must not make tech_stack data depend on overview data
    # existing too (see the unconditional call site in build_report).
    prs = _prs()
    slide = add_solutions_products_slide(prs, {}, {"hostname": "example.com", "ip": "1.2.3.4", "https": True})
    assert slide is not None
    texts = _texts(slide)
    assert "Tech Stack & Hosting" in texts
    assert "Solutions" not in texts and "Products (by category)" not in texts


def test_title_is_renamed():
    prs = _prs()
    slide = add_solutions_products_slide(prs, {"solutions": ["A solution"]}, {})
    title = next(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    assert title == "Solutions, Products & Tech Stack"


def test_two_cards_at_spec_geometry():
    prs = _prs()
    slide = add_solutions_products_slide(prs, {"solutions": ["A solution"]}, {"hostname": "x.com"})
    # Left + right card only — filters out the full-slide background,
    # the top bar (full width, Inches(0.1) tall) and hosting-row divider
    # lines (wide but Pt(0.75) tall) by bounding both width and height to
    # card scale (cards are 7.8in/4.0in wide, 5.9in tall).
    cards = [
        sh for sh in slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and Inches(3) < sh.width < Inches(10) and sh.height > Inches(1)
    ]
    assert len(cards) == 2
    left_card, right_card = sorted(cards, key=lambda c: c.left)
    assert left_card.width > right_card.width  # left card (7.8in) wider than right (4.0in)


def test_cms_pulled_out_of_tile_grid_and_normalized():
    prs = _prs()
    tech = {
        "hostname": "x.com", "https": True,
        "detected": [{"category": "cms", "name": "wordpress"}, {"category": "cdn", "name": "Cloudflare"}],
    }
    slide = add_solutions_products_slide(prs, {}, tech)
    texts = _texts(slide)
    assert "Wordpress" in texts or "WordPress" in texts  # title-cased, never raw lowercase
    assert texts.count("Cloudflare") == 1  # shown once, in the tile grid — not duplicated as a hosting row
    assert "CMS" in texts


def test_https_colored_and_missing_ip_shows_em_dash():
    prs = _prs()
    tech = {"hostname": "insecure.example", "ip": None, "https": False, "detected": []}
    slide = add_solutions_products_slide(prs, {}, tech)
    texts = _texts(slide)
    assert "✗ Not secure" in texts
    assert "✓ Secure" not in texts
    assert "—" in texts  # missing IP address row
    assert "No technologies detected" in texts


def test_tile_names_shortened_via_known_aliases():
    prs = _prs()
    tech = {
        "hostname": "x.com", "https": True,
        "detected": [
            {"category": "analytics", "name": "Google Analytics (GA4)"},
            {"category": "cdn", "name": "Amazon CloudFront"},
        ],
    }
    slide = add_solutions_products_slide(prs, {}, tech)
    texts = _texts(slide)
    assert "Google Analytics 4" in texts
    assert "CloudFront" in texts
    assert "Google Analytics (GA4)" not in texts
    assert "Amazon CloudFront" not in texts


def test_more_than_six_technologies_shows_overflow_tile():
    prs = _prs()
    detected = [{"category": "other", "name": f"Tool {i}"} for i in range(9)]
    tech = {"hostname": "x.com", "https": True, "detected": detected}
    slide = add_solutions_products_slide(prs, {}, tech)
    texts = _texts(slide)
    assert "+4 more" in texts  # 9 items, 6 slots -> 5 real tiles shown + 1 counter tile covering the other 4
    assert sum(1 for t in texts if t.startswith("Tool ")) == 5


def test_products_heading_never_renders_as_an_orphan_with_no_content():
    # Regression (confirmed real, QA stress case, 2026-09-25): Solutions
    # can legitimately consume nearly the whole left card on a client with
    # many/long solution bullets, and drawing the "Products (by category)"
    # heading unconditionally then left it with nothing rendered under it.
    overview = {
        "solutions": [f"A fairly long solution description number {i} describing a real capability in detail" for i in range(10)],
        "products_by_category": {f"Category {i}": [f"Product {i}-{j}" for j in range(6)] for i in range(8)},
    }
    prs = _prs()
    slide = add_solutions_products_slide(prs, overview, {})
    texts = _texts(slide)
    if "Products (by category)" in texts:
        idx = texts.index("Products (by category)")
        assert idx + 1 < len(texts)  # something real follows the heading, never the last line on the slide
    assert _audit_slide_geometry(prs) == []


def test_solutions_truncate_with_plus_n_more_when_they_dont_all_fit():
    # Each item long enough to wrap to 2 lines, and enough of them that
    # even at ~0.30in/line the full list can't fit in the card.
    overview = {"solutions": [
        f"Solution number {i} with a genuinely long description that will wrap across two lines in this column width, describing a real capability in detail"
        for i in range(20)
    ]}
    prs = _prs()
    slide = add_solutions_products_slide(prs, overview, {})
    texts = _texts(slide)
    assert any(t.startswith("•  +") and "more" in t for t in texts)
    assert _audit_slide_geometry(prs) == []


def test_long_product_descriptions_shrink_font_before_dropping_content():
    overview = {
        "products_by_category": {
            f"Category {i}": [f"Product {i}-{j} with a fairly long descriptive name" for j in range(5)]
            for i in range(4)
        },
    }
    prs = _prs()
    slide = add_solutions_products_slide(prs, overview, {})
    assert _audit_slide_geometry(prs) == []
    texts = _texts(slide)
    assert "Category 0" in texts  # at least the first category survives


def test_no_overlap_or_overflow_across_all_three_qa_cases():
    # Mirrors the spec's own 3 required QA cases.
    case_a = (
        {"solutions": ["24/7 managed DBA support", "Database modernization"],
         "products_by_category": {"Database Services": ["MySQL", "PostgreSQL", "Oracle"]}},
        {"hostname": "www.example.com", "ip": "104.21.45.12", "https": True,
         "detected": [{"category": "cms", "name": "WordPress"}, {"category": "cdn", "name": "Cloudflare"}]},
    )
    case_b = (
        {"solutions": [f"Solution {i} with a good amount of descriptive text in it" for i in range(8)],
         "products_by_category": {f"Category {i}": [f"Product {i}-{j} long name" for j in range(5)] for i in range(6)}},
        {"hostname": "www.example-long-hostname.co.uk", "ip": "192.168.1.100", "https": True,
         "detected": [{"category": "other", "name": f"Tool {i}"} for i in range(7)]},
    )
    case_c = (
        {"solutions": ["SEO audits"], "products_by_category": {"Services": ["Audit"]}},
        {"hostname": "insecuresite.example", "ip": None, "https": False, "detected": [{"category": "cms", "name": "WordPress"}]},
    )
    for overview, tech in (case_a, case_b, case_c):
        prs = _prs()
        add_solutions_products_slide(prs, overview, tech)
        assert _audit_slide_geometry(prs) == []

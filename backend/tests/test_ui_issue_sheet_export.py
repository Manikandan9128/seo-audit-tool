from unittest.mock import MagicMock, patch

from app.services.ui_issue_sheet_export import create_ui_issues_sheet, export_full_issue_list, export_ui_issues_xlsx

_MOD = "app.services.ui_issue_sheet_export"


def _issues():
    return [
        {"title": "No clear CTA", "evidence": "0 CTAs above the fold", "where": "Hero", "fix": "Add a CTA", "priority": "High", "device": "Both"},
        {"title": "Small tap target", "evidence": "32px tall button", "where": "Mobile header", "fix": "Increase to 48px", "priority": "Low", "device": "Mobile"},
    ]


# --- create_ui_issues_sheet ---------------------------------------------------

def test_create_sheet_returns_none_when_no_issues():
    assert create_ui_issues_sheet("Acme", [], db=None) is None


def test_create_sheet_returns_none_when_not_connected():
    with patch(f"{_MOD}.get_sheets_oauth_credentials", return_value=None):
        result = create_ui_issues_sheet("Acme", _issues(), db=MagicMock())
    assert result is None


def test_create_sheet_returns_edit_url_on_success():
    with patch(f"{_MOD}.get_sheets_oauth_credentials", return_value=MagicMock()), \
         patch(f"{_MOD}.build") as mock_build, \
         patch(f"{_MOD}._create_spreadsheet_file", return_value="sheet123") as mock_create_file:
        mock_sheets = MagicMock()
        mock_drive = MagicMock()
        mock_build.side_effect = lambda name, version, credentials: {"sheets": mock_sheets, "drive": mock_drive}[name]
        result = create_ui_issues_sheet("Acme", _issues(), db=MagicMock())
    assert result == "https://docs.google.com/spreadsheets/d/sheet123/edit"
    mock_create_file.assert_called_once()
    # The values written include the header + one row per issue.
    update_call = mock_sheets.spreadsheets.return_value.values.return_value.update
    written_values = update_call.call_args.kwargs["body"]["values"]
    assert written_values[0] == ["#", "Issue", "Evidence", "Where", "Fix", "Priority", "Device"]
    assert written_values[1][1] == "No clear CTA"
    assert written_values[2][1] == "Small tap target"
    mock_drive.permissions.return_value.create.assert_called_once()


def test_create_sheet_returns_none_and_disconnects_on_dead_refresh_token():
    from google.auth.exceptions import RefreshError

    with patch(f"{_MOD}.get_sheets_oauth_credentials", side_effect=RefreshError("invalid_grant")), \
         patch(f"{_MOD}.disconnect_sheets_oauth") as mock_disconnect:
        result = create_ui_issues_sheet("Acme", _issues(), db=MagicMock())
    assert result is None
    mock_disconnect.assert_called_once()


def test_create_sheet_returns_none_on_any_other_api_failure():
    with patch(f"{_MOD}.get_sheets_oauth_credentials", return_value=MagicMock()), \
         patch(f"{_MOD}.build", side_effect=RuntimeError("Sheets API down")):
        result = create_ui_issues_sheet("Acme", _issues(), db=MagicMock())
    assert result is None


# --- export_ui_issues_xlsx ----------------------------------------------------

def test_export_xlsx_contains_header_and_all_rows():
    from openpyxl import load_workbook

    xlsx_bytes = export_ui_issues_xlsx(_issues())
    wb = load_workbook(filename=__import__("io").BytesIO(xlsx_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("#", "Issue", "Evidence", "Where", "Fix", "Priority", "Device")
    assert rows[1] == (1, "No clear CTA", "0 CTAs above the fold", "Hero", "Add a CTA", "High", "Both")
    assert rows[2][0] == 2
    assert len(rows) == 3


# --- export_full_issue_list ---------------------------------------------------

def test_export_full_issue_list_returns_nothing_when_no_issues():
    result = export_full_issue_list("Acme", [], db=None)
    assert result == {"full_list_url": None, "xlsx_bytes": None}


def test_export_full_issue_list_prefers_sheets_when_available():
    with patch(f"{_MOD}.create_ui_issues_sheet", return_value="https://sheet.url"):
        result = export_full_issue_list("Acme", _issues(), db=MagicMock())
    assert result == {"full_list_url": "https://sheet.url", "xlsx_bytes": None}


def test_export_full_issue_list_falls_back_to_xlsx_when_sheets_unavailable():
    with patch(f"{_MOD}.create_ui_issues_sheet", return_value=None):
        result = export_full_issue_list("Acme", _issues(), db=MagicMock())
    assert result["full_list_url"] is None
    assert result["xlsx_bytes"] is not None
    assert isinstance(result["xlsx_bytes"], bytes)

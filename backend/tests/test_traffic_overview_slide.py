"""Traffic Overview slide (add_traffic_overview_slide) — Traffic Sources
channel table geometry regression.

Confirmed real (Geopits regen, 2026-09-26): the channel table's col_widths
([2.9, 1.4, 1.4], summing to 5.7in) didn't match the `width` passed to
add_table (3.3in) — _draw_table renders each column at its own col_widths
entry regardless of `width`, so the table's real right edge landed at
8.8in, deep inside the Key Insights column that starts at x=6.7in. Text
boxes were never checked against tables for overlap, so this was invisible
to _audit_slide_geometry until it was extended to include TABLE/CHART
shapes the same day."""

from pptx import Presentation

from app.reporting import pptx_builder
from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _audit_slide_geometry, add_traffic_overview_slide


def _prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def _analytics():
    rows = [{
        "sessions": 3735, "total_users": 3322, "page_views": 5940,
        "engagement_duration": 26800, "engagement_rate": 0.268, "bounce_rate": 0.732,
    }]
    sources = [
        {"channel": "Direct", "sessions": 2540},
        {"channel": "Organic Search", "sessions": 934},
        {"channel": "Paid Social", "sessions": 120},
        {"channel": "Organic Social", "sessions": 78},
        {"channel": "Referral", "sessions": 41},
        {"channel": "Email", "sessions": 22},
    ]
    return {
        "date_range": {"ga4_start": "2026-08-27", "ga4_end": "2026-09-25"},
        "traffic_overview": {"rows": rows},
        "traffic_sources": {"rows": sources},
    }


def test_traffic_sources_table_never_overlaps_key_insights():
    prs = _prs()
    pptx_builder._theme["footer"] = "Geopits  ·  www.geopits.com"
    try:
        add_traffic_overview_slide(prs, _analytics())
        issues = _audit_slide_geometry(prs)
        assert issues == [], issues
    finally:
        pptx_builder._theme["footer"] = ""


def test_traffic_sources_table_right_edge_stays_left_of_key_insights_column():
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu, Inches

    prs = _prs()
    slide = add_traffic_overview_slide(prs, _analytics())
    table = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.TABLE)
    key_insights_label = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and sh.text_frame.text.strip() == "KEY INSIGHTS"
    )
    assert table.left + table.width <= key_insights_label.left, (
        f"table right edge ({Emu(table.left + table.width).inches:.2f}in) "
        f"overlaps the Key Insights column (starts {Emu(key_insights_label.left).inches:.2f}in)"
    )

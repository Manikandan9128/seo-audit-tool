"""Universal SEO keyword engine — strategy layer (2026-09-24): funnel,
audience, secondary intent, time-sensitive keywords (§8/§10/§11/§44),
§31 opportunity score, §66-K roadmap, §18 hierarchy, §38 coverage, §37
internal links, §56 page map, and their slides / Sheet columns."""

from pptx import Presentation

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, _audit_slide_geometry, add_content_seo_next_steps_slide, add_keyword_topic_map_slide,
)
from app.services.google_sheets_service import _CLIENT_COLUMNS, _CLIENT_HEADER
from app.services.keyword_intelligence_service import (
    detect_intent, funnel_stage, is_temporal, keyword_audience, secondary_intent,
)
from app.services.keyword_strategy_service import (
    OPPORTUNITY_WEIGHTS, apply_strategy_to_manual_clusters, build_keyword_strategy, keyword_opportunity,
    summaries_from_keyword_rows,
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                parts.extend(c.text_frame.text for c in row.cells)
    return "\n".join(parts)


def _rows(kws, family="Commercial", relevance=None):
    return [{"keyword": k, "search_volume": v, "keyword_difficulty": 25, "detected_intent": family,
             "intent_family": family, "relevance_status": relevance} for k, v in kws]


def _summary(name, kws, family="Commercial", url=None, strength="none", relevance=None, confidence="High"):
    return {"name": name, "rows": _rows(kws, family, relevance), "target_url": url, "match_strength": strength,
            "confidence_level": confidence, "dominant_family": family}


# --- per-keyword attributes (§8, §10, §11, §44) -----------------------------

def test_funnel_stage_audience_secondary_intent_and_time_sensitivity():
    assert funnel_stage("truck overheating fix", "Commercial") == "Problem Discovery"
    assert funnel_stage("how to choose running shoes", "Informational") == "Education"
    assert funnel_stage("tipper truck", "Commercial") == "Solution Discovery"
    assert funnel_stage("acme customer care", "Navigational") == "Retention / Support"
    assert keyword_audience("payroll software for small business") == "Business / Enterprise"
    assert keyword_audience("educational toys for 2 year olds") == "Parent / Child"
    assert keyword_audience("tipper truck") is None  # never guessed
    primary = detect_intent("best truck price")["intent"]
    assert secondary_intent("best truck price", primary) == "Commercial Investigation"
    assert is_temporal("bus price 2026") and is_temporal("latest tipper models") and not is_temporal("tipper truck")


# --- §31 opportunity + §66-K roadmap ----------------------------------------

def test_opportunity_is_not_volume_alone():
    relevant_small = {"keyword": "a", "search_volume": 500, "keyword_difficulty": 20,
                      "detected_intent": "Transactional", "relevance_status": "Core Relevant"}
    flagged_big = {"keyword": "b", "search_volume": 50000, "keyword_difficulty": 20,
                   "detected_intent": "Commercial", "relevance_status": "Product/Service Mismatch"}
    assert keyword_opportunity(relevant_small, 50000)["score"] > keyword_opportunity(flagged_big, 50000)["score"]


def test_missing_evidence_factors_are_left_out_not_guessed():
    result = keyword_opportunity({"keyword": "a", "search_volume": 100}, 100)
    assert result["factors"]["serp_opportunity"] is None
    assert result["factors"]["competitor_gap"] is None
    assert 0 <= result["score"] <= 100
    assert abs(sum(OPPORTUNITY_WEIGHTS.values()) - 1.0) < 1e-9


def test_roadmap_sends_mostly_flagged_or_low_confidence_clusters_to_review():
    strategy = build_keyword_strategy([
        _summary("Buses", [("school bus", 22200), ("bus chassis", 2900)]),
        _summary("Price", [("brabus price", 2400), ("peterbilt price", 390)], relevance="Product/Service Mismatch"),
        _summary("Odd", [("odd thing", 50)], confidence="Low"),
    ])
    assert "Price" in strategy["roadmap"]["Human Review"] and "Odd" in strategy["roadmap"]["Human Review"]
    assert "Buses" in strategy["roadmap"]["High"] + strategy["roadmap"]["Medium"]


# --- §18 / §38 / §37 topic model ---------------------------------------------

def test_topic_hierarchy_groups_clusters_under_their_parent_entity():
    strategy = build_keyword_strategy([
        _summary("Trucks", [("truck", 110000), ("trucks", 4000)]),
        _summary("Truck Types", [("tipper truck", 14800), ("tanker truck", 3600)]),
        _summary("Truck Parts", [("truck parts", 880), ("truck spare parts", 880)], url="https://bb.com/parts", strength="strong"),
        _summary("Truck Specs", [("truck sizes", 390)], family="Informational"),
        _summary("Buses", [("school bus", 22200)], url="https://bb.com/buses", strength="partial"),
    ])
    topics = {t["parent"]: t for t in strategy["topics"]}
    assert set(topics["Truck"]["clusters"]) == {"Trucks", "Truck Types", "Truck Parts", "Truck Specs"}
    assert topics["Truck"]["coverage"] == "Weak" and topics["Truck"]["covered"] == 1
    assert "Informational" in topics["Truck"]["intents_without_page"]
    assert topics["Bus"]["coverage"] == "Strong"
    links = {(l["from_cluster"], l["to_cluster"]): l["relation"] for l in topics["Truck"]["links"]}
    assert links[("Truck Types", "Trucks")] == "Subtopic → Parent hub"
    assert links[("Truck Specs", "Trucks")] == "Guide → Product/Service"


def test_over_concentrated_token_does_not_become_one_giant_parent():
    # Regression (confirmed real, Geopits report, 2026-09-24): "Database"
    # held 100 clusters, "SQL" held 92 — every cluster shared "database"
    # so it always won as parent, turning the topic map into a couple of
    # rows reading "+97 more" that told the reader nothing. 16 clusters
    # here all share "database" but each also has its own distinct second
    # word — they should split onto their own more specific parents
    # instead of piling under one "Database" bucket.
    words = ["support", "managed", "backup", "migration", "performance", "security", "monitoring",
             "consulting", "tuning", "replication", "sharding", "partitioning", "indexing",
             "clustering", "scaling", "hosting"]
    strategy = build_keyword_strategy([
        _summary(f"Database {w.title()}", [(f"database {w}", 100)]) for w in words
    ])
    parents = {t["parent"] for t in strategy["topics"]}
    assert not any(len(t["clusters"]) > 15 for t in strategy["topics"])
    assert len(parents) > 1  # never all 16 collapsed into one "Database" bucket


def test_weakest_topic_insight_prefers_real_demand_over_first_in_list():
    # Regression (confirmed real, Geopits report, 2026-09-24): "Weakest
    # topic: Trigger" (a ~200-combined-search topic) was recommended to
    # build first, ahead of topics carrying vastly more real demand, just
    # because it happened to be first among Weak/Missing topics in
    # opportunity-sort order.
    strategy = build_keyword_strategy([
        _summary("Trigger Types", [("trigger types", 110), ("list triggers", 90)]),
        _summary("Database Support", [("database support services", 260), ("remote database support", 140)]),
    ])
    slide = add_keyword_topic_map_slide(_prs(), strategy)
    text = _text(slide)
    # "Database Support" (400 combined searches) must win over "Trigger
    # Types" (200) — whichever parent token it resolves under (a 2-word
    # cluster's own parent-label tie-break is a separate concern from
    # this fix), never the lower-demand topic.
    assert "Weakest topic:" in text
    assert "Weakest topic: Trigger" not in text


def test_sibling_products_are_not_linked_just_for_sharing_a_word():
    strategy = build_keyword_strategy([
        _summary("Truck Types", [("tipper truck", 14800)]),
        _summary("Truck Parts", [("truck parts", 880)]),
    ])
    assert strategy["topics"][0]["links"] == []


# --- §56 page map --------------------------------------------------------------

def test_page_map_flags_two_clusters_on_one_url():
    strategy = build_keyword_strategy([
        _summary("A", [("alpha widget", 900)], url="https://x.com/w", strength="strong"),
        _summary("B", [("beta widget", 800)], url="https://x.com/w", strength="partial"),
        _summary("C", [("gamma gadget", 700)]),
    ])
    pages = {p["url"]: p for p in strategy["page_map"]}
    assert pages["https://x.com/w"]["cannibalization_risk"] and set(pages["https://x.com/w"]["clusters"]) == {"A", "B"}
    assert pages[None]["clusters"] == ["C"] and not pages[None]["cannibalization_risk"]


# --- both paths feed it --------------------------------------------------------

def test_routing_buckets_are_never_pages():
    rows = [
        {"keyword": "tipper truck", "cluster": "Tippers", "search_volume": 100, "cluster_source": "rule"},
        {"keyword": "truck driver jobs", "cluster": "Jobs / Careers", "search_volume": 900, "cluster_source": "routing"},
    ]
    assert [s["name"] for s in summaries_from_keyword_rows(rows)] == ["Tippers"]


def test_manual_clusters_get_opportunity_and_priority_without_touching_the_sheet():
    clusters = [{"cluster": "Buses", "keywords": [
        {"keyword": "school bus", "search_volume": 22200, "intent": "Commercial"},
        {"keyword": "bus chassis", "search_volume": 2900, "intent": "Commercial"},
    ], "target_url": None, "match_strength": "none", "confidence_level": "High"}]
    strategy = apply_strategy_to_manual_clusters(clusters)
    assert clusters[0]["roadmap_priority"] in ("High", "Medium", "Low") and clusters[0]["opportunity"] > 0
    assert set(clusters[0]["keywords"][0]) == {"keyword", "search_volume", "intent"}
    assert strategy["topics"][0]["parent"] == "Bus"


# --- slides / sheet -------------------------------------------------------------

def test_topic_map_slide_renders_and_fits():
    summaries = [
        _summary(f"Very Long Descriptive Cluster Name {i}", [(f"widget model{i} alpha", 1000 - i), (f"widget model{i} beta", 500)],
                 family="Informational" if i % 3 == 0 else "Commercial")
        for i in range(12)
    ]
    summaries.append(_summary("Widgets", [("widget", 9000)]))
    strategy = build_keyword_strategy(summaries)
    prs = _prs()
    slide = add_keyword_topic_map_slide(prs, strategy)
    text = _text(slide)
    assert "Keyword Topic Map" in text and "Widget" in text and "Internal link" in text
    assert _audit_slide_geometry(prs) == []
    assert add_keyword_topic_map_slide(_prs(), None) is None


def test_content_seo_next_steps_follow_the_roadmap_not_volume():
    rows = [
        {"keyword": "big doubtful", "cluster": "Doubtful", "search_volume": 90000, "relevance_status": "Relevant",
         "roadmap_priority": "Human Review", "cluster_opportunity": 20, "page_category": "Landing Page",
         "existing_page_action": "Create New Page"},
        {"keyword": "small strong", "cluster": "Strong", "search_volume": 300, "relevance_status": "Relevant",
         "roadmap_priority": "High", "cluster_opportunity": 80, "page_category": "Landing Page",
         "existing_page_action": "Create New Page", "primary_or_secondary": "Primary"},
    ]
    text = _text(add_content_seo_next_steps_slide(_prs(), rows))
    assert text.index('"Strong" topic') < text.index('"Doubtful" topic')
    assert "[High priority, opportunity 80/100]" in text and "[Needs human review" in text
    assert 'primary keyword "small strong"' in text


def test_sheet_client_tab_is_the_master_dataset_and_keeps_original_columns_first():
    assert _CLIENT_HEADER[:5] == ["Keyword", "Cluster", "Search Volume", "KD", "Intent"]
    keys = {key for _label, key in _CLIENT_COLUMNS}
    for field in ("funnel_stage", "audience", "secondary_intent", "opportunity_score", "roadmap_priority",
                  "user_need", "recommended_action"):
        assert field in keys


def test_no_internal_link_from_a_page_to_itself():
    strategy = build_keyword_strategy([
        _summary("Certified Payroll", [("certified payroll", 4400)], url="https://x.com/cp", strength="strong"),
        _summary("Certified Payroll Guides", [("what is certified payroll", 1000)], family="Informational",
                 url="https://x.com/cp", strength="strong"),
    ])
    assert strategy["topics"][0]["links"] == []
    assert strategy["page_map"][0]["cannibalization_risk"]

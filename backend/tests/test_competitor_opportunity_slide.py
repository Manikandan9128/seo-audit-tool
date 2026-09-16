import base64

from pptx import Presentation
from pptx.util import Inches

from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, add_competitor_opportunity_slide, add_competitor_opportunity_summary_slide,
)

# Minimal valid 1x1 transparent PNG — a well-known tiny test fixture, used
# here to prove add_picture actually succeeds (not just that a bad-bytes
# except-and-skip path was hit).
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    for s in slide.shapes:
        if s.has_table:
            for row in s.table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def test_slide_absent_when_nothing_to_show():
    assert add_competitor_opportunity_slide(_prs(), "Client", "rival.com", {}) is None


def test_slide_renders_headline_unique_angle_and_gap():
    narrative = {
        "headline": "usage-based pricing calculator",
        "unique_angle": ["Lets visitors price a plan by seat count with no sales call."],
        "gap": "Client has no self-serve pricing tool of its own.",
    }
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    text = _slide_text(slide)
    assert "rival.com's Unique Angle: usage-based pricing calculator" in text
    assert "UNIQUE ANGLE" in text
    assert "Lets visitors price a plan by seat count" in text
    assert "GAP FOR CLIENT" in text
    assert "no self-serve pricing tool" in text
    assert "SHARED ADVANTAGE" not in text  # omitted entirely when not shared


def test_shared_advantage_only_rendered_when_present():
    narrative = {
        "headline": "usage-based pricing calculator",
        "unique_angle": ["mechanic"],
        "gap": "gap text",
        "shared_advantage": "Also offered by miter.com — not unique to rival.com.",
    }
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    text = _slide_text(slide)
    assert "SHARED ADVANTAGE" in text
    assert "Also offered by miter.com" in text


def test_screenshot_position_and_size_unchanged():
    narrative = {"headline": "h", "unique_angle": ["a"], "gap": "g", "screenshot": _TINY_PNG}
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    pictures = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
    assert len(pictures) == 1
    pic = pictures[0]
    assert pic.left == Inches(8.7)
    assert pic.width == Inches(3.9)
    assert pic.top == Inches(1.1)  # card_top, unchanged constant


def test_long_bullets_never_overlap_or_run_past_the_card():
    # Standing rule (2026-09-11): content must always fit the slide, never
    # overlap. Feed deliberately long text and check every text box's
    # vertical span stays within the card and none overlap another.
    narrative = {
        "headline": "a very long unique angle headline that will wrap across more than one line on the slide",
        "unique_angle": [
            "A long, specific, mechanics-heavy bullet describing exactly how this competitor's pricing calculator works step by step. " * 2,
        ],
        "gap": "An equally long gap bullet explaining precisely what the client lacks in direct response to that mechanism. " * 2,
        "shared_advantage": "A long shared-advantage explanation naming the other competitor this overlaps with in detail. " * 2,
    }
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    spans = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            spans.append((shape.top, shape.top + shape.height))
    spans.sort()
    for (top1, bottom1), (top2, _) in zip(spans, spans[1:]):
        assert bottom1 <= top2 + Inches(0.01)  # small tolerance for rounding
    card_bottom = Inches(1.1) + (Inches(7.15) - Inches(1.1))
    assert all(bottom <= card_bottom + Inches(0.05) for _, bottom in spans)


def test_full_schema_renders_all_new_sections():
    # 2026-09-16 user spec: Evidence / Why It Matters / Opportunity /
    # Implementation, in addition to the existing Unique Angle / Gap.
    narrative = {
        "headline": "usage-based pricing calculator",
        "unique_angle": ["Lets visitors price a plan by seat count with no sales call."],
        "evidence": [
            "Ranks #3 for \"team pricing calculator\" (2,400 searches/mo).",
            "14 of 22 tracked keywords are pricing/comparison-shaped.",
        ],
        "why_it_matters": "This captures decision-stage buyers who are actively comparing vendors on price.",
        "gap": "Client has no self-serve pricing tool of its own.",
        "opportunity": "Build an interactive seat-based pricing calculator on the pricing page.",
        "implementation": [
            "Ship a pricing calculator widget on /pricing.",
            "Target the same comparison-shaped keyword cluster in on-page copy.",
        ],
    }
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    text = _slide_text(slide)
    assert "EVIDENCE" in text
    assert "Ranks #3 for" in text
    assert "WHY IT MATTERS" in text
    assert "decision-stage buyers" in text
    assert "OPPORTUNITY" in text
    assert "interactive seat-based pricing calculator" in text
    assert "IMPLEMENTATION" in text
    assert "pricing calculator widget" in text


def test_full_schema_with_long_text_never_overlaps_or_runs_past_card():
    narrative = {
        "headline": "a very long unique angle headline that will wrap across more than one line on the slide",
        "unique_angle": ["A long mechanics-heavy bullet describing the pricing calculator step by step. " * 2],
        "evidence": [
            "A long evidence bullet citing several specific keywords, positions and volumes in detail. " * 2,
            "A second long evidence bullet citing backlink and referring-domain figures in detail. " * 2,
        ],
        "why_it_matters": "A long why-it-matters paragraph explaining the strategic importance in detail. " * 3,
        "gap": "An equally long gap bullet explaining precisely what the client lacks in direct response. " * 2,
        "opportunity": "A long opportunity sentence naming the specific initiative the client should build. " * 2,
        "implementation": [
            "A long first implementation step with specific detail about execution. " * 2,
            "A long second implementation step with specific detail about execution. " * 2,
        ],
        "shared_advantage": "A long shared-advantage explanation naming the other competitor in detail. " * 2,
    }
    slide = add_competitor_opportunity_slide(_prs(), "Client", "rival.com", narrative)
    spans = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            spans.append((shape.top, shape.top + shape.height))
    spans.sort()
    for (top1, bottom1), (top2, _) in zip(spans, spans[1:]):
        assert bottom1 <= top2 + Inches(0.01)
    card_bottom = Inches(1.1) + (Inches(7.15) - Inches(1.1))
    assert all(bottom <= card_bottom + Inches(0.05) for _, bottom in spans)


def test_summary_slide_absent_when_nothing_to_show():
    assert add_competitor_opportunity_summary_slide(_prs(), {}, None) is None
    assert add_competitor_opportunity_summary_slide(_prs(), {"rival.com": {"error": "failed"}}, None) is None


def test_summary_slide_renders_table_and_top_opportunities():
    narratives = {
        "rival.com": {"headline": "pricing calculator", "gap": "no self-serve pricing tool", "opportunity": "build a calculator"},
        "other.com": {"headline": "comparison hub", "gap": "no vs-pages", "opportunity": "build comparison pages"},
    }
    top_opportunities = [
        "Build a self-serve pricing calculator to capture decision-stage, price-comparing buyers.",
        "Launch a comparison-page hub targeting head-to-head competitor searches.",
    ]
    slide = add_competitor_opportunity_summary_slide(_prs(), narratives, top_opportunities)
    text = _slide_text(slide)
    assert "Cross-Competitor Opportunity Summary" in text
    assert "rival.com" in text and "other.com" in text
    assert "pricing calculator" in text
    assert "Build a self-serve pricing calculator" in text

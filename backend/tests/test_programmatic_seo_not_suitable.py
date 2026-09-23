from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_programmatic_seo_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_states_not_suitable_when_every_cluster_fails_demand_gate():
    # 2026-09-21 spec section 11's own required wording: "No validated
    # programmatic SEO opportunity identified..." must be an explicit
    # stated finding, not a silently-omitted slide, when real clusters
    # existed but none cleared eligibility.
    rows = [
        {"keyword": "small topic a", "search_volume": 10, "cluster": "Small Topic"},
        {"keyword": "small topic b", "search_volume": 10, "cluster": "Small Topic"},
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    assert slide is not None
    text = _slide_text(slide)
    assert "No validated programmatic SEO opportunity identified" in text
    assert "search volume" in text.lower()


def test_states_not_suitable_when_no_cluster_has_enough_distinct_subpages():
    rows = [
        {"keyword": "widget", "search_volume": 500, "cluster": "Widget"},
        {"keyword": "widget", "search_volume": 400, "cluster": "Widget"},  # duplicate keyword, no new sub-intent
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    assert slide is not None
    text = _slide_text(slide)
    assert "No validated programmatic SEO opportunity identified" in text


def test_returns_none_when_there_is_no_clustered_keyword_data_at_all():
    # No real clusters at all (not even a failed candidate) — nothing was
    # evaluated, so this stays a silent omission, not a false "not
    # suitable" claim about data that was never assessed.
    rows = [{"keyword": "kw", "search_volume": 500}]  # no "cluster" key
    assert add_programmatic_seo_slide(_prs(), rows) is None
    assert add_programmatic_seo_slide(_prs(), None) is None


def test_renders_real_opportunities_when_eligible():
    rows = [
        {"keyword": "widget pricing calculator", "search_volume": 1000, "cluster": "Widgets"},
        {"keyword": "widget installation guide", "search_volume": 900, "cluster": "Widgets"},
        {"keyword": "widget maintenance tips", "search_volume": 800, "cluster": "Widgets"},
        {"keyword": "widget troubleshooting steps", "search_volume": 700, "cluster": "Widgets"},
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "No validated" not in text
    assert "Widgets" in text
    assert "New hub + 4 subpages" in text
    assert "3.4K/mo" in text  # combined volume, real number not a placeholder


def test_no_backend_debug_terminology_ever_reaches_the_ppt():
    # 2026-09-21 spec section 1 — these internal validation labels drove
    # the exclusion logic but must never appear in the rendered slide.
    rows = [
        {"keyword": "daimler", "search_volume": 500, "cluster": "Dealership"},
        {"keyword": "daimler india", "search_volume": 400, "cluster": "Dealership"},
        {"keyword": "daimler india commercial vehicles pvt ltd", "search_volume": 300, "cluster": "Dealership"},
        {"keyword": "bharat benz showroom", "search_volume": 900, "cluster": "Dealership"},
        {"keyword": "bharat benz showroom near me", "search_volume": 800, "cluster": "Dealership"},
        {"keyword": "bharatbenz near me", "search_volume": 700, "cluster": "Dealership"},
        {"keyword": "benz dealership contact", "search_volume": 600, "cluster": "Dealership"},
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    text = _slide_text(slide)
    for banned in (
        "NEAR_DUPLICATE_PAIR", "SUBPAGE_DUPLICATES_HUB", "CLUSTER_DEBUG", "INTERNAL_VALIDATION",
        "EXCLUDED_KEYWORD", "SOURCE_CLUSTER", "MODEL_REASON", "CANDIDATE_CLUSTER", "Data-quality note",
    ):
        assert banned not in text


def test_near_duplicates_consolidated_not_shown_as_separate_subpages():
    # "Daimler" / "Daimler India" / "Daimler India Commercial Vehicles Pvt
    # Ltd" are the same underlying intent — must collapse to ONE subpage,
    # not three, and the why-it-qualifies line should say so in plain
    # language (never the internal NEAR_DUPLICATE_PAIR label).
    rows = [
        {"keyword": "daimler", "search_volume": 500, "cluster": "Dealership"},
        {"keyword": "daimler india", "search_volume": 400, "cluster": "Dealership"},
        {"keyword": "daimler india commercial vehicles pvt ltd", "search_volume": 300, "cluster": "Dealership"},
        {"keyword": "bharat benz showroom", "search_volume": 900, "cluster": "Dealership"},
        {"keyword": "bharat benz showroom near me", "search_volume": 800, "cluster": "Dealership"},
        {"keyword": "bharatbenz near me", "search_volume": 700, "cluster": "Dealership"},
        {"keyword": "benz dealership contact", "search_volume": 600, "cluster": "Dealership"},
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    text = _slide_text(slide)
    assert "near-duplicate variants folded in" in text


def test_ranking_prefers_more_distinct_subintents_over_raw_volume():
    # 2026-09-21 spec section 12: never rank by search volume alone. A
    # lower-volume cluster with 4 genuinely distinct sub-intents must
    # outrank a higher-volume cluster that only clears the 3-subpage floor.
    rows = [
        {"keyword": "alpha pricing", "search_volume": 5000, "cluster": "Alpha Topic"},
        {"keyword": "alpha guide", "search_volume": 4000, "cluster": "Alpha Topic"},
        {"keyword": "alpha support", "search_volume": 3000, "cluster": "Alpha Topic"},
        {"keyword": "beta pricing", "search_volume": 500, "cluster": "Beta Topic"},
        {"keyword": "beta guide", "search_volume": 400, "cluster": "Beta Topic"},
        {"keyword": "beta support", "search_volume": 300, "cluster": "Beta Topic"},
        {"keyword": "beta reviews", "search_volume": 200, "cluster": "Beta Topic"},
    ]
    slide = add_programmatic_seo_slide(_prs(), rows)
    table = next(s for s in slide.shapes if s.has_table).table
    first_data_row = [c.text_frame.text for c in table.rows[1].cells]
    assert first_data_row[0] == "Beta Topic"  # 4 sub-intents beats Alpha's higher volume with only 3


def test_caps_at_max_opportunities_never_manufactures_extra():
    distinct_subs = ["pricing calculator", "installation guide", "maintenance tips", "troubleshooting steps"]
    rows = []
    for cluster_i in range(8):
        for sub_i, sub in enumerate(distinct_subs):
            rows.append({
                "keyword": f"topic{cluster_i} {sub}",
                "search_volume": 1000 - cluster_i * 10 - sub_i,
                "cluster": f"Topic {cluster_i}",
            })
    slide = add_programmatic_seo_slide(_prs(), rows)
    table = next(s for s in slide.shapes if s.has_table).table
    assert len(table.rows) - 1 == 5  # header + at most 5 opportunity rows


# --- Programmatic SEO spec 2026-09-23 --------------------------------------
_WIDGETS = [
    {"keyword": "widget pricing calculator", "search_volume": 1000, "cluster": "Widgets"},
    {"keyword": "widget installation guide", "search_volume": 900, "cluster": "Widgets"},
    {"keyword": "widget maintenance tips", "search_volume": 800, "cluster": "Widgets"},
    {"keyword": "widget troubleshooting steps", "search_volume": 700, "cluster": "Widgets"},
]


def test_routing_bucket_clusters_never_become_programmatic_opportunities():
    rows = [dict(r, cluster="Competitor / Comparison Opportunities") for r in _WIDGETS]
    assert add_programmatic_seo_slide(_prs(), rows) is None


def test_existing_hub_page_is_reused_not_recreated():
    rows = [dict(r, existing_page_url="https://x.com//widgets", existing_page_match_strength="strong") for r in _WIDGETS]
    text = _slide_text(add_programmatic_seo_slide(_prs(), rows))
    assert "Existing hub + 4 subpages" in text and "existing hub page https://x.com/widgets" in text


def test_insights_are_distinct_per_opportunity():
    rows = list(_WIDGETS) + [
        {"keyword": "gadget repair cost", "search_volume": 900, "cluster": "Gadgets"},
        {"keyword": "gadget setup manual", "search_volume": 800, "cluster": "Gadgets"},
        {"keyword": "gadget battery life", "search_volume": 700, "cluster": "Gadgets"},
    ]
    text = _slide_text(add_programmatic_seo_slide(_prs(), rows))
    assert "Widgets: 4 distinct sub-intents" in text and "Gadgets: 3 distinct sub-intents" in text

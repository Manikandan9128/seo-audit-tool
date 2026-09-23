from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_aeo_geo_visibility_required_slide
from app.services.geopulse_ai_service import _drop_cross_list_duplicates, _item_violation


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    return "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())


RAW = "Workflow / How-To: 0/10 unbranded mentions. Cited source: https://example.com/guides/fleet"


def test_no_visibility_check_states_the_gap_instead_of_generic_advice():
    text = _slide_text(add_aeo_geo_visibility_required_slide(_prs()))
    assert "AI visibility check required" in text
    assert "ChatGPT" not in text and "schema" not in text.lower()


def test_uploaded_but_unavailable_is_worded_differently():
    text = _slide_text(add_aeo_geo_visibility_required_slide(_prs(), uploaded_but_unavailable=True))
    assert "was uploaded" in text and "required" not in text.split("\n")[1]


def test_generic_schema_cause_and_invented_url_items_are_dropped():
    assert _item_violation("Create more content.", RAW) == "generic recommendation"
    assert _item_violation("Add FAQPage schema to the fleet guides.", RAW).startswith("schema")
    assert _item_violation("The brand was not mentioned because its pages are thin.", RAW) == "assumed cause"
    assert _item_violation("Expand https://example.com/pricing with fleet sizing answers.", RAW).startswith("URL not in")
    ok = "Create fleet-operation guides covering tipper capacity and servicing to address the 0/10 Workflow / How-To gap; see https://example.com/guides/fleet."
    assert _item_violation(ok, RAW) is None


def test_geo_item_restating_an_aeo_item_is_dropped():
    aeo = ["Create fleet-operation guides covering tipper capacity, fuel efficiency and servicing for the Workflow gap."]
    geo = [
        "Create fleet-operation guides covering tipper capacity, fuel efficiency and servicing for the Workflow gap.",
        "Strengthen third-party citations around commercial-vehicle topics — no client-domain sources were cited in this run.",
    ]
    assert _drop_cross_list_duplicates(aeo, geo) == [geo[1]]

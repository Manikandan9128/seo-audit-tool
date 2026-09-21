from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_pagespeed_score_breakdown_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def _mobile(**overrides):
    base = {
        "current_score": 32,
        "metric_table": [
            {"id": "largest-contentful-paint", "label": "LCP", "value": 18800, "display_value": "18.8 s", "good_threshold": 2500, "status": "Poor"},
            {"id": "cumulative-layout-shift", "label": "CLS", "value": 0.42, "display_value": "0.42", "good_threshold": 0.1, "status": "Poor"},
        ],
        "field_data": None,
        "opportunities": [],
        "diagnostics": [],
        "lcp_breakdown": None,
    }
    base.update(overrides)
    return base


def test_returns_empty_list_when_no_metric_table():
    assert add_pagespeed_score_breakdown_slide(_prs(), None, None) == []
    assert add_pagespeed_score_breakdown_slide(_prs(), {"metric_table": []}, None) == []


def test_never_contains_banned_projection_language():
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(), None)
    text = "\n".join(_slide_text(s) for s in slides)
    for banned in ("If Fixed", "Score Impact", "score_if_fixed", "score_delta"):
        assert banned not in text


def test_metric_table_renders_current_threshold_and_status_columns():
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(), None)
    text = _slide_text(slides[0])
    assert "18.8 s" in text
    assert "2.5s" in text  # good threshold, formatted
    assert "Poor" in text


def test_missing_field_data_states_insufficient_evidence():
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(field_data=None), None)
    text = _slide_text(slides[0])
    assert "Insufficient field data for Core Web Vitals assessment." in text


def test_field_data_present_shows_real_categories_not_inferred():
    field_data = {
        "is_origin_fallback": False, "overall_category": "SLOW",
        "lcp": {"percentile": 4200, "category": "SLOW"},
        "inp": None,
        "cls": {"percentile": 8, "category": "AVERAGE"},
    }
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(field_data=field_data), None)
    text = _slide_text(slides[0])
    assert "LCP: SLOW" in text
    assert "INP: no field data" in text


def test_diagnosis_only_for_non_good_metrics():
    mobile = _mobile(metric_table=[
        {"id": "largest-contentful-paint", "label": "LCP", "value": 2000, "display_value": "2.0 s", "good_threshold": 2500, "status": "Good"},
        {"id": "total-blocking-time", "label": "TBT", "value": 900, "display_value": "900 ms", "good_threshold": 200, "status": "Poor"},
    ])
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, mobile, None)
    text = _slide_text(slides[0])
    assert "TBT indicates significant main-thread activity" in text
    assert "LCP is a performance constraint" not in text


def test_opportunities_slide_only_added_when_data_present():
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(), None)
    assert len(slides) == 1  # no opportunities/diagnostics/lcp_breakdown supplied

    prs2 = _prs()
    slides2 = add_pagespeed_score_breakdown_slide(prs2, _mobile(opportunities=[{"title": "Eliminate render-blocking resources", "savings": "1.2s", "savings_ms": 1200}]), None)
    assert len(slides2) == 2
    text = _slide_text(slides2[1])
    assert "Eliminate render-blocking resources" in text
    assert "1.2s" in text


def test_opportunities_slide_never_invents_a_recommendation():
    prs = _prs()
    slides = add_pagespeed_score_breakdown_slide(prs, _mobile(diagnostics=[{"id": "bootup-time", "label": "Total JavaScript execution time", "value": "2.1 s"}]), None)
    text = _slide_text(slides[1])
    for banned in ("Compress images", "Remove unused JavaScript", "Optimize fonts", "Improve internal linking"):
        assert banned not in text
    assert "Total JavaScript execution time: 2.1 s" in text

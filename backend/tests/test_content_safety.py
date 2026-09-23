"""Hard rule (2026-09-23): no 18+ content anywhere in a client report."""

from types import SimpleNamespace

from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _table_slide
from app.services.content_safety import SafeImport, is_adult, redact_presentation, scrub


def test_detects_adult_terms_including_concatenated_and_leetspeak_forms():
    for kw in [
        "indian bus xxx", "xnxx bus", "xxxvideo", "pornhub", "desi sex video", "p0rn", "s3x video",
        "call girls near me", "18+ video", "nude photos", "onlyfans leaks", "sexvideos",
        "school bus bf", "indian bus flash",
    ]:
        assert is_adult(kw), kw


def test_legitimate_business_keywords_are_not_blocked():
    for kw in [
        "sexual harassment policy", "escort vehicle for oversize load", "xxl t shirt", "xxxl hoodie",
        "unisex uniforms", "essex payroll services", "school bus", "construction payroll", "tipper truck",
        "bf goodrich truck tyres", "flash tipper price",
    ]:
        assert not is_adult(kw), kw


def test_scrub_drops_adult_rows_and_keys_and_sentences_without_mutating_input():
    data = {
        "keyword_rows": [{"keyword": "school bus"}, {"keyword": "indian bus xxx"}],
        "analytics": {"search_queries": {"rows": [{"query": "xnxx bus"}, {"query": "bharatbenz"}]}},
        "detected": {"indian bus xxx": "Commercial", "school bus": "Commercial"},
        "core_problem": {"summary": "Gap is large. Top gap: indian bus xxx at 2,900/mo. Fix titles."},
        "insights": ["Top: \"indian bus xxx\"", "School bus demand is high."],
    }
    out = scrub(data)
    assert [r["keyword"] for r in out["keyword_rows"]] == ["school bus"]
    assert [r["query"] for r in out["analytics"]["search_queries"]["rows"]] == ["bharatbenz"]
    assert list(out["detected"]) == ["school bus"]
    assert "xxx" not in out["core_problem"]["summary"] and "Fix titles." in out["core_problem"]["summary"]
    assert out["insights"] == ["School bus demand is high."]
    assert len(data["keyword_rows"]) == 2  # input untouched


def test_safe_import_never_modifies_stored_upload_data():
    stored = {"rows": [{"keyword": "school bus"}, {"keyword": "indian bus xxx"}], "row_count": 2}
    imp = SimpleNamespace(import_type="keyword_gap", is_own_site=True, parsed_data=stored)
    view = SafeImport(imp)
    assert [r["keyword"] for r in view.parsed_data["rows"]] == ["school bus"]
    assert view.import_type == "keyword_gap" and view.is_own_site
    assert len(stored["rows"]) == 2  # DB object untouched (additive-only rule)


def test_final_deck_safety_net_removes_adult_table_rows_and_text():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _table_slide(prs, "Target Keywords", ["Keyword", "Volume"], [("school bus", "22,200"), ("indian bus xxx", "2,900")],
                 insights=['Highest demand: "indian bus xxx" — 2,900/mo.', "School bus leads demand."])
    removed = redact_presentation(prs)
    text = "\n".join(
        (sh.text_frame.text if sh.has_text_frame else "\n".join(c.text_frame.text for r in sh.table.rows for c in r.cells))
        for sh in prs.slides[0].shapes if sh.has_text_frame or sh.has_table
    )
    assert removed >= 2
    assert "xxx" not in text
    assert "school bus" in text

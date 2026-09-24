"""Keyword engine quality fixes from the Geopits no-file report review
(2026-09-24). Every rule here is universal — the Geopits-shaped data is
only the example that exposed it."""

from unittest.mock import patch

from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_keyword_research_slide, add_keyword_topic_map_slide
from app.services.keyword_cluster_pipeline import build_final_keyword_clusters
from app.services.keyword_intelligence_service import (
    annotate_keyword_rows, build_rule_groups, detect_intent, display_phrase, rule_group_name, smart_title,
)
from app.services.keyword_relevance_service import (
    _rule_exclude, is_junk_keyword, vendor_product_pricing_reason, vendor_tokens,
)
from app.services.keyword_strategy_service import build_keyword_strategy, summaries_from_keyword_rows

_AI_OFF = (
    patch("app.services.keyword_cluster_pipeline.generate_business_themes", return_value={}),
    patch("app.services.keyword_cluster_pipeline.generate_phase2_candidate_clusters", return_value=({}, [])),
    patch("app.services.keyword_cluster_pipeline.generate_phase3_validated_clusters", return_value=[]),
)


def _run(rows, pages=None):
    with _AI_OFF[0], _AI_OFF[1], _AI_OFF[2]:
        build_final_keyword_clusters(rows, "Acme", None, pages)
    by = {}
    for r in rows:
        by.setdefault(r["cluster"], []).append(r["keyword"])
    return by


def _rows(items, intent="Commercial", category="Landing Page"):
    return [{"keyword": k, "search_volume": v, "keyword_difficulty": 30, "intent": intent,
             "page_category": category, "relevance_status": "Relevant"} for k, v in items]


# 1 — §39/§41/§43 same need on the same page is one cluster --------------------

def test_same_need_clusters_on_one_page_are_merged():
    rows = _rows([("sql index", 720), ("index in sql", 390), ("database indexing", 720), ("indexing in database", 260),
                  ("sql server", 5000), ("sql query", 300), ("database backup", 200), ("sql join", 100)],
                 intent="Informational", category="Blog / Guide")
    pages = [{"page_url": "https://x.com/blog/what-is-an-index-in-sql-database-index",
              "page_title": "What is an Index in SQL and How a Database Index Works"}]
    by = _run(rows, pages)
    index_cluster = next(k for k, v in by.items() if "sql index" in v)
    assert {"sql index", "database indexing", "indexing in database"} <= set(by[index_cluster])


def test_clusters_sharing_only_a_site_wide_head_word_are_not_merged():
    rows = _rows([("sql index", 720), ("sql backup", 700), ("sql join", 300), ("sql server", 900), ("sql view", 200)],
                 intent="Informational", category="Blog / Guide")
    pages = [{"page_url": "https://x.com/blog/sql", "page_title": "SQL Index Backup Join Server View"}]
    by = _run(rows, pages)
    assert not any({"sql index", "sql backup"} <= set(v) for v in by.values())


# 2 — §40 a head term with many different subtopics is a category ------------

def test_category_head_does_not_swallow_its_subtopics():
    rows = [{"keyword": k, "search_volume": 100} for k in [
        "sql server", "sql server etl", "sql server encryption", "sql server replication", "sql server backup",
        "sql server monitoring", "sql server health"]]
    annotate_keyword_rows(rows)
    groups = build_rule_groups(rows)
    assert max(len(g["rows"]) for g in groups) == 1


def test_small_parent_still_absorbs_its_few_variants():
    rows = [{"keyword": k, "search_volume": 100} for k in [
        "certified payroll", "government certified payroll", "certified payroll software"]]
    annotate_keyword_rows(rows)
    assert len(build_rule_groups(rows)) == 1


# 3 — §21 junk is never a keyword ----------------------------------------------

def test_hash_url_and_code_fragments_are_junk_but_real_queries_are_not():
    assert is_junk_keyword("7c3735cb431b03e60e736b893c594322ba391c6f cipherdetails create sql")
    assert is_junk_keyword("www.example.com") and is_junk_keyword("select * from t where a=1")
    for real in ("a/b testing", "24/7 dba support", "bs6 engine", "iphone 15 pro max", "1015r truck", "sql server 2019"):
        assert not is_junk_keyword(real), real
    assert _rule_exclude("7c3735cb431b03e60e736b893c594322ba391c6f cipherdetails", set())


def test_junk_keyword_never_becomes_a_cluster():
    rows = _rows([("tipper truck", 1000), ("7c3735cb431b03e60e736b893c594322ba391c6f cipherdetails create sql", 20)])
    by = _run(rows)
    assert not any("7c37" in name.lower() or "Cipherdetails" in name for name in by)


# 4 — §55/§60 readable names ---------------------------------------------------

def test_cluster_names_keep_acronyms_and_never_repeat_words():
    assert smart_title("sql server server") == "SQL Server"
    assert smart_title(display_phrase("postgresql vs mysql")) == "PostgreSQL vs MySQL"
    assert smart_title(display_phrase("what is azure data studio")) == "Azure Data Studio"
    assert smart_title(display_phrase("aws aurora pricing")) == "AWS Aurora"
    rows = [{"keyword": "postgresql vs mysql", "search_volume": 100}]
    annotate_keyword_rows(rows)
    assert rule_group_name(build_rule_groups(rows)[0]) == "PostgreSQL vs MySQL"


# 5 — §8/§46 intent -----------------------------------------------------------

def test_either_or_is_a_comparison_and_troubleshooting_is_a_guide():
    assert detect_intent("mysql or sql server")["intent"] == "Comparison"
    assert detect_intent("to be or not to be")["intent"] != "Comparison"
    assert detect_intent("sql server error 18456")["family"] == "Informational"
    assert detect_intent("cost based optimizer")["intent"] != "Transactional"
    assert detect_intent("truck price")["intent"] == "Transactional"


# 6 — §18 topic map: filler words are never topics, insights keep room -----------

def test_filler_word_is_never_a_parent_topic_and_rows_are_compact():
    rows = _rows([("cost based optimizer", 3600), ("rule based optimizer", 200)]
                 + [(f"widget model{i}", 500 - i) for i in range(30)])
    _run(rows)
    strategy = build_keyword_strategy(summaries_from_keyword_rows(rows))
    assert all(t["parent"].lower() != "based" for t in strategy["topics"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = add_keyword_topic_map_slide(prs, strategy)
    table = next(sh.table for sh in slide.shapes if sh.has_table)
    assert len(table.rows) <= 7  # header + 6 topics
    assert all(len(row.cells[1].text.split(", ")) <= 3 for row in list(table.rows)[1:])


# 7 — §66-K slides follow the roadmap -------------------------------------------

def test_target_keyword_slides_follow_roadmap_priority():
    rows = []
    for name, tier, opp in (("Low One", "Low", 30), ("High One", "High", 80), ("Medium One", "Medium", 55)):
        for kw in (f"{name.lower()} a", f"{name.lower()} b"):
            rows.append({"keyword": kw, "cluster": name, "search_volume": 1000 if tier == "Low" else 10,
                         "roadmap_priority": tier, "cluster_opportunity": opp})
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    add_keyword_research_slide(prs, rows)
    titles = [next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
              for s in prs.slides]
    assert titles == ["Target Keywords: High One", "Target Keywords: Medium One", "Target Keywords: Low One"]


# 8 — §21 another company's product pricing --------------------------------------

def test_vendor_product_pricing_is_flagged_but_client_services_are_not():
    pages = [{"page_url": "https://x.com/partners/aws"}, {"page_url": "https://x.com/technologies/oracle-database"},
             {"page_url": "https://x.com/services/remote-dba"}]
    vendors = vendor_tokens(pages, {"acme"})
    assert {"aws", "oracle"} <= vendors
    assert vendor_product_pricing_reason("aws aurora pricing", vendors)
    assert not vendor_product_pricing_reason("oracle dba services cost", vendors)
    assert not vendor_product_pricing_reason("aws migration pricing", vendors)
    assert not vendor_product_pricing_reason("database pricing", vendors)
    assert not vendor_product_pricing_reason("aws aurora pricing", set())


def test_vendor_pricing_rows_go_to_review_not_to_a_target_slide():
    rows = _rows([("aws aurora pricing", 390), ("remote dba services", 1300)])
    rows[0]["relevance_status"] = "Unknown / Needs Review"
    rows[0]["relevance_reason"] = "Price/purchase search for another company's product"
    by = _run(rows)
    review = next(name for name, kws in by.items() if "aws aurora pricing" in kws)
    assert review.startswith("Needs Review")

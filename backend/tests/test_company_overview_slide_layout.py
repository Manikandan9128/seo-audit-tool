"""Company Overview slide (add_company_overview_extracted_slide) layout
regression: the footer (registration/contact) line must never land on top
of the left column's real content."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.reporting import pptx_builder
from app.reporting.pptx_builder import (
    SLIDE_H, SLIDE_W, _audit_slide_geometry, _chip_row, add_company_overview_extracted_slide,
)


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def test_footer_never_overlaps_wrapped_industry_chips():
    # Regression (confirmed real, Geopits report, 2026-09-24): the left
    # card's height was a fixed constant, not measured from real content.
    # A long description + 6 KPIs + industries wrapping onto a second chip
    # row pushed real content past that fixed height, but the footer
    # (registration/contact) still drew at the fixed top+height offset —
    # landing directly on top of the wrapped "E-Learning / EdTech" /
    # "Retail" chips instead of below them.
    overview = {
        "company_name": "Geopits",
        "description": (
            "Geopits is a database management and data engineering services firm specializing in "
            "24/7 managed DBA support, database modernization, and cloud migration. The agency "
            "delivers end-to-end data platform engineering, performance tuning, high availability "
            "architecture, and generative AI readiness solutions across hybrid and cloud environments."
        ),
        "kpis": [
            "200+ Clients Globally", "2000+ Databases Managed", "140+ Successful Migrations Completed",
            "25% cloud cost reduction", "99.99% average pipeline uptime", "15-minute response for P1 critical outages",
        ],
        "industries": ["Automotive", "Non-Banking Financial Companies (NBFC)", "E-Learning / EdTech", "Retail"],
        "target_country": "Global",
        "target_market": "Mid-Market and Enterprise",
        "primary_buyers": ["CTO", "CIO", "VP of Engineering", "Head of Data", "Director of Infrastructure"],
        "daily_users": ["Database Administrators", "Data Engineers", "DevOps Engineers", "Software Developers"],
        "beneficiaries": ["In-house Engineering Teams", "Data Analytics Teams", "Executives", "Business Operations"],
        "registration_info": "ISO 27001 Certified, ISO 9001 Certified",
    }
    prs = _prs()
    slide = add_company_overview_extracted_slide(prs, "Geopits", overview)

    chip_bottoms = [
        sh.top + sh.height for sh in slide.shapes
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.has_text_frame
        and sh.text_frame.text.strip() in overview["industries"]
    ]
    assert chip_bottoms, "expected industry chips to be found on the slide"

    footer = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and "ISO 27001 Certified" in sh.text_frame.text
    )
    assert footer.top >= max(chip_bottoms), (
        f"footer (top={footer.top}) overlaps the industry chips (bottom={max(chip_bottoms)})"
    )


def test_footer_never_renders_past_the_bottom_of_the_slide():
    # Regression (confirmed real, Geopits regen, 2026-09-25): flooring the
    # footer on whichever column ran longest fixed the overlap above, but
    # with enough ICP items the RIGHT column itself can wrap past the
    # slide's own bottom edge — the footer then floors on that and runs
    # off the visible page entirely (bottom=7.64in on a 7.50in slide).
    overview = {
        "company_name": "Geopits",
        "description": "Short description.",
        "kpis": ["KPI one", "KPI two"],
        "industries": ["Automotive", "Retail"],
        "target_country": "Global",
        "target_market": "Enterprise",
        "primary_buyers": [f"Buyer role {i}" for i in range(6)],
        "daily_users": [f"Daily user role number {i}" for i in range(6)],
        "beneficiaries": [f"Beneficiary group with a fairly long name {i}" for i in range(6)],
        "registration_info": "ISO 27001 Certified, ISO 9001 Certified",
    }
    prs = _prs()
    slide = add_company_overview_extracted_slide(prs, "Geopits", overview)
    footer = next(
        sh for sh in slide.shapes
        if sh.has_text_frame and "ISO 27001 Certified" in sh.text_frame.text
    )
    assert footer.top + footer.height <= SLIDE_H, (
        f"footer bottom ({footer.top + footer.height}) runs past the slide height ({SLIDE_H})"
    )


def test_registration_line_never_overlaps_the_persistent_global_footer():
    # Regression (confirmed real, Geopits regen, 2026-09-26): both prior
    # tests here call add_company_overview_extracted_slide() directly
    # without ever setting _theme["footer"], so _footer() (pptx_builder.py)
    # no-ops and neither test could ever see this. In a real build_report()
    # run, build_report() sets _theme["footer"] = "{client} · {domain}"
    # BEFORE any slide function runs, and _content_header() draws it on
    # every content slide at a fixed row (SLIDE_H-0.4in to SLIDE_H-0.1in,
    # L=0.4-8.4in). The registration/contact line's old bottom-of-page cap
    # (SLIDE_H-0.42in) put it on that exact same row — a guaranteed
    # collision on every deck with a company_overview payload, not an edge
    # case: "Geopits · www.geopits.com" <-> "ISO 27001 Certified, ISO 9001
    # Certified".
    pptx_builder._theme["footer"] = "Geopits  ·  www.geopits.com"
    try:
        overview = {
            "company_name": "Geopits",
            "description": "Short description.",
            "kpis": ["KPI one", "KPI two"],
            "industries": ["Automotive", "Retail"],
            "registration_info": "ISO 27001 Certified, ISO 9001 Certified",
        }
        prs = _prs()
        add_company_overview_extracted_slide(prs, "Geopits", overview)
        issues = _audit_slide_geometry(prs)
        assert not any("overlap" in i for i in issues), issues
    finally:
        pptx_builder._theme["footer"] = ""


def test_industry_chip_boxes_are_wide_enough_for_their_bold_text():
    # Regression (confirmed real, Geopits regen 2026-09-26): "E-Learning /
    # EdTech" and "Retail" both visibly overflowed their own pill — the
    # per-char width estimate (~0.075in-per-char-at-size-11, i.e. the
    # REGULAR-weight rate) was never adjusted for the fact every chip's
    # text renders bold (run.font.bold = True). 0.62in-per-72pt-per-char is
    # the calibrated bold-text constant this codebase already trusts
    # elsewhere (_ui_fixes_truncate_cell) — every chip box must be at least
    # that wide for its own text plus padding.
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    size = 10.5
    items = ["Automotive", "Non-Banking Financial Companies (NBFC)", "E-Learning / EdTech", "Retail"]
    _chip_row(slide, Inches(0.6), Inches(1.0), items, Inches(6.0), size=size)

    pad_x = Inches(0.14)
    min_char_w = Emu(int(Inches(1) * 0.62 * size / 72))
    for item in items:
        chip = next(
            sh for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip() == item
        )
        required = Emu(int(min_char_w) * len(item)) + pad_x * 2
        assert chip.width >= required - Emu(1000), (  # 1000 EMU (~0.001in) float-rounding slack
            f"{item!r} chip is {Emu(chip.width).inches:.3f}in wide, needs >= {Emu(required).inches:.3f}in for bold text"
        )

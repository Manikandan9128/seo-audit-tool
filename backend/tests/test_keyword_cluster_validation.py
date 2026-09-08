from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_keyword_research_slide


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_mixed_page_formats_flags_split():
    rows = [
        {"keyword": "what is certified payroll", "cluster": "C", "search_volume": 500, "page_category": "Blog / Guide"},
        {"keyword": "how does certified payroll work", "cluster": "C", "search_volume": 300, "page_category": "Blog / Guide"},
        {"keyword": "certified payroll software", "cluster": "C", "search_volume": 900, "page_category": "Landing Page"},
        {"keyword": "certified payroll pricing", "cluster": "C", "search_volume": 200, "page_category": "Landing Page"},
    ]
    slides = add_keyword_research_slide(_prs(), rows)
    text = _slide_text(slides[0])
    assert "mixed page formats" in text
    assert "consider splitting" in text


def test_consistent_page_format_keeps_cluster():
    rows = [
        {"keyword": "certified payroll software", "cluster": "C", "search_volume": 900, "page_category": "Landing Page"},
        {"keyword": "certified payroll pricing", "cluster": "C", "search_volume": 200, "page_category": "Landing Page"},
    ]
    slides = add_keyword_research_slide(_prs(), rows)
    text = _slide_text(slides[0])
    assert "consistent format (Landing Page)" in text
    assert "mixed page formats" not in text


def test_single_outlier_keyword_does_not_trigger_split():
    # Regression guard: one stray keyword in a different format shouldn't
    # flag a whole cluster for splitting — needs a real second sub-group.
    rows = [
        {"keyword": "certified payroll software", "cluster": "C", "search_volume": 900, "page_category": "Landing Page"},
        {"keyword": "random outlier keyword", "cluster": "C", "search_volume": 100, "page_category": "Blog / Guide"},
    ]
    slides = add_keyword_research_slide(_prs(), rows)
    text = _slide_text(slides[0])
    assert "mixed page formats" not in text
    assert "consistent format" in text


def test_no_page_category_data_omits_validation_bullet():
    rows = [{"keyword": "some keyword", "cluster": "C", "search_volume": 500}]
    slides = add_keyword_research_slide(_prs(), rows)
    text = _slide_text(slides[0])
    assert "Cluster validation" not in text

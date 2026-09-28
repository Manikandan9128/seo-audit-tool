from unittest.mock import patch

from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, _audit_slide_geometry, add_core_problem_slide
from app.services.core_problem_service import CORE_PROBLEM_PROMPT, generate_core_problem


def _attempts(*items):
    """items: list of (raw_text, provider) tuples to yield in order."""
    def _gen(prompt, max_tokens, errors):
        for raw, provider in items:
            yield raw, provider
    return _gen


def test_degenerate_null_response_returns_error_not_crash():
    # 2026-09-15 live crash: Groq returning the literal token "null" (valid
    # JSON, parses to None with no exception) hit data.get("thesis") on a
    # NoneType and surfaced as "'NoneType' object has no attribute 'get'"
    # during report generation instead of falling through to the error path.
    with patch("app.services.core_problem_service.iter_text_attempts", side_effect=_attempts(("null", "groq"))):
        result = generate_core_problem({"issues": []})
    assert "error" in result


def test_well_formed_response_still_parses():
    raw = '{"thesis": "Site lacks structured data.", "categories": []}'
    with patch("app.services.core_problem_service.iter_text_attempts", side_effect=_attempts((raw, "groq"))):
        result = generate_core_problem({"issues": []})
    assert result["thesis"] == "Site lacks structured data."


def test_first_provider_invalid_json_falls_through_to_next():
    # 2026-09-22: OpenRouter's auto-router can land on a model that
    # doesn't reliably follow "return ONLY valid JSON" — a bad response
    # from one provider must not kill the section outright, it should
    # fall through to whichever provider comes next.
    good = '{"thesis": "Site lacks structured data.", "categories": []}'
    with patch(
        "app.services.core_problem_service.iter_text_attempts",
        side_effect=_attempts(("not json at all", "openrouter"), (good, "gemini")),
    ):
        result = generate_core_problem({"issues": []})
    assert result["thesis"] == "Site lacks structured data."


def test_prompt_forbids_adding_the_off_topic_excluded_count_back_into_a_total():
    # Regression (confirmed real, Geopits regen 2026-09-26): Core Problem's
    # slide stated "315 relevant keyword gaps (304 missing)" while the
    # Competitor Keyword Gap slides — built from the SAME authoritative
    # gap_counts, per build_keyword_gap_summary_finding's own single-
    # source-of-truth contract — said "301 (290 missing)". 315-301=14 and
    # 304-290=14: the model read the finding's own "(14 off-topic/
    # competitor-brand excluded)" parenthetical and added it back into
    # both the total and the missing subcount instead of treating it as
    # already-excluded context. This only reaches the model via free-text
    # prompt instruction (there's no code path that assembles Core
    # Problem's numbers directly), so the regression this pins is the
    # instruction's presence, not the model's actual compliance.
    assert "never be added back into it" in CORE_PROBLEM_PROMPT
    assert "off-topic/competitor-brand excluded" in CORE_PROBLEM_PROMPT

    captured = {}

    def _gen(prompt, max_tokens, errors):
        captured["prompt"] = prompt
        yield '{"thesis": "t", "categories": []}', "groq"

    finding = {
        "summary": "301 relevant keyword gap(s) (11 shared, 290 missing), 12,000 combined monthly searches "
                    "(14 off-topic/competitor-brand excluded)",
    }
    with patch("app.services.core_problem_service.iter_text_attempts", side_effect=_gen):
        generate_core_problem({"competitor_gap_findings": [finding]})
    assert "301 relevant keyword gap" in captured["prompt"]
    assert "14 off-topic/competitor-brand excluded" in captured["prompt"]


def test_prompt_asks_for_flat_2_to_3_problems_not_per_category():
    # 2026-09-28 redesign: the slide's own final-implementation rule forbids
    # the old 3-category/up-to-12-finding format ("show only the 2-3 most
    # consequential confirmed problems") — pin that the prompt actually asks
    # for a flat ranked list, not one mandatory pick per topic area.
    assert '"problems": [string]' in CORE_PROBLEM_PROMPT
    assert "categories" not in CORE_PROBLEM_PROMPT
    assert "2 or 3 items" in CORE_PROBLEM_PROMPT


def test_parse_caps_problems_at_three_and_drops_internal_field_leaks():
    raw = ('{"thesis": "t", "problems": ["Real finding one.", "backlink_summary returned null", '
           '"Real finding two.", "Real finding three.", "Real finding four."]}')

    def _gen(prompt, max_tokens, errors):
        yield raw, "groq"

    with patch("app.services.core_problem_service.iter_text_attempts", side_effect=_gen):
        result = generate_core_problem({"issues": []})
    assert result["problems"] == ["Real finding one.", "Real finding two.", "Real finding three."]


def test_core_problem_slide_renders_flat_ranked_list_with_no_geometry_violations():
    core_problem = {
        "thesis": "The site's organic visibility is currently constrained by thin content and weak backlink authority.",
        "problems": [
            "The audit identified 47 technical errors across 823 crawled pages, concentrated on category pages.",
            "Organic visibility is currently constrained by a Domain Rating of 18, well below the tracked competitor average of 42.",
        ],
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = add_core_problem_slide(prs, core_problem)
    assert slide is not None
    assert _audit_slide_geometry(prs) == []


def test_core_problem_slide_shows_three_when_three_are_given_never_pads():
    core_problem = {
        "thesis": "Thesis.",
        "problems": ["Problem one is here with a full sentence of real detail.",
                     "Problem two is here with a full sentence of real detail.",
                     "Problem three is here with a full sentence of real detail."],
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = add_core_problem_slide(prs, core_problem)
    assert _audit_slide_geometry(prs) == []
    text_boxes = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith(("1.", "2.", "3."))]
    assert len(text_boxes) == 3

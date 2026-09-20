from pptx import Presentation

from app.reporting.pptx_builder import SLIDE_H, SLIDE_W, add_ux_findings_slides


def _prs():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        elif shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text_frame.text)
    return "\n".join(parts)


def test_column_header_is_principle_not_bias():
    ux_findings = {
        "onboarding_breakdown": [
            {"principle": "Risk Reversal", "where": "Primary CTA", "suggestion": "Add reassurance copy near the CTA."},
        ],
    }
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    text = _slide_text(onboarding_slide)
    assert "Principle / Heuristic" in text
    assert "\nBias\n" not in text


def test_accepts_legacy_bias_key_for_backward_compat():
    ux_findings = {"onboarding_breakdown": [{"bias": "Social Proof", "where": "Hero", "suggestion": "Add proof points."}]}
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    assert "Social Proof" in _slide_text(onboarding_slide)


def test_summary_insight_has_no_conversion_impact_ranking_claim():
    ux_findings = {
        "onboarding_breakdown": [
            {"principle": "Cognitive Load", "where": "Nav", "suggestion": "Reduce choices."},
            {"principle": "Risk Reversal", "where": "CTA", "suggestion": "Add guarantee."},
        ],
    }
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    text = _slide_text(onboarding_slide)
    assert "ranked by likely impact" not in text
    assert "sign-up/purchase completion" not in text
    assert "Cognitive Load" in text and "Risk Reversal" in text


def test_source_reflects_vision_pass_when_flagged():
    ux_findings = {
        "onboarding_breakdown": [{"principle": "Framing", "where": "Banner", "suggestion": "Clarify value prop."}],
        "onboarding_breakdown_source": "vision",
    }
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    assert "Homepage screenshot analysis" in _slide_text(onboarding_slide)


def test_source_reflects_manual_notes_when_not_flagged_as_vision():
    ux_findings = {"onboarding_breakdown": [{"principle": "Framing", "where": "Banner", "suggestion": "Clarify value prop."}]}
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    assert "Manual UX walkthrough" in _slide_text(onboarding_slide)


def test_item_count_is_not_padded_to_five():
    ux_findings = {"onboarding_breakdown": [{"principle": "Framing", "where": "Banner", "suggestion": "Clarify value prop."}]}
    slides = add_ux_findings_slides(_prs(), ux_findings)
    onboarding_slide = next(s for s in slides if "Onboarding Breakdown" in _slide_text(s))
    table_shape = next(s for s in onboarding_slide.shapes if s.has_table)
    assert len(table_shape.table.rows) == 2  # header + 1 real row, no padding
